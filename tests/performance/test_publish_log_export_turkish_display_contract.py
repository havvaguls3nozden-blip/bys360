"""BYS360 performance publish-log Excel export -- raw DB action_type leak
closure (H1D).

Root cause (found during the H1D repo-wide scan, same defect class as the
Phase 8 "checkpoint" action_type leak fixed in H1B):
PerformancePublishLog.action_type is stored as a raw English code
("publish", "unpublish", "bulk_publish", "bulk_unpublish" -- see the model
comment at app/models/performance_models.py:653) and was written verbatim
into the "İşlem Türü" (Action Type) column of the
/performance/publish/logs/export/excel admin Excel export
(app/performance/engagement_publish_routes.py).

Fix: a new humanize_publish_log_action()/PUBLISH_LOG_ACTION_LABELS pair was
added to app/services/performance_admin_service.py (co-located with
create_publish_log(), which writes these values) and is now used at the one
place these values reach a human (the Excel export). The stored
PerformancePublishLog.action_type column and the log_action query filter
are completely untouched -- this is presentation-only.

Writes NOTHING to any production source file -- only to its own isolated,
temporary SQLite DB.
"""
from __future__ import annotations

import io
import tempfile
import uuid
from datetime import date
from pathlib import Path

import pytest

_TEST_DB_ROOT = Path(tempfile.gettempdir()) / "bys360" / "publish_log_export_turkish_tmp" / "test_dbs"
_PASSWORD = "PublishLogExportTurkishTest1!"


def _make_app(monkeypatch):
    _TEST_DB_ROOT.mkdir(parents=True, exist_ok=True)
    db_path = _TEST_DB_ROOT / f"{uuid.uuid4().hex}.sqlite3"

    monkeypatch.setenv("APP_ENV", "testing")
    monkeypatch.setenv("FLASK_ENV", "testing")
    monkeypatch.setenv("SECRET_KEY", "test-secret-key-for-publish-log-export-turkish-contract")
    monkeypatch.setenv("DEFAULT_FIRST_LOGIN_PASSWORD", "publish-log-export-first-login-test-pw")
    monkeypatch.setenv("FLASK_SKIP_SCHEMA_VALIDATION", "1")
    monkeypatch.setenv("AUTO_REPAIR_SCHEMA", "false")
    monkeypatch.setenv("STRICT_SCHEMA_CHECK", "false")
    monkeypatch.setenv("REQUIRE_DOTENV_FILE", "false")
    monkeypatch.setenv("STRICT_ENV_VALIDATION", "false")
    monkeypatch.setenv("WTF_CSRF_ENABLED", "false")
    monkeypatch.setenv("DATABASE_URL", "sqlite:///" + db_path.as_posix())
    monkeypatch.setenv("MAIL_SUPPRESS_SEND", "true")
    monkeypatch.setenv("SCHEDULER_ENABLED", "false")
    monkeypatch.setenv("LOGIN_FORCE_CAPTCHA_FOR_UNKNOWN_USER", "false")

    from app import create_app
    from config import Config

    monkeypatch.setattr(Config, "APP_ENV", "testing")
    monkeypatch.setattr(Config, "SQLALCHEMY_DATABASE_URI", "sqlite:///" + db_path.as_posix())
    monkeypatch.setattr(Config, "SQLALCHEMY_ENGINE_OPTIONS", {})

    app = create_app()
    app.config.update(TESTING=True, WTF_CSRF_ENABLED=False, SQLALCHEMY_DATABASE_URI="sqlite:///" + db_path.as_posix())

    from app.extensions import db

    with app.app_context():
        db.create_all()

    return app


@pytest.fixture
def app(monkeypatch):
    return _make_app(monkeypatch)


@pytest.fixture
def client(app):
    return app.test_client()


def _create_user(app, *, sicil_no, role="admin", password=_PASSWORD):
    from app.extensions import db
    from app.models import User, UserMenuPermission

    with app.app_context():
        user = User(
            sicil_no=sicil_no,
            email=f"{sicil_no}@ktb.gov.tr",
            ad="Publish",
            soyad="LogExportContract",
            role=role,
            is_active=True,
            must_change_password=False,
            must_set_security_question=False,
        )
        user.set_password(password)
        db.session.add(user)
        db.session.commit()
        db.session.add(UserMenuPermission(user_id=user.id, menu_key="performance_publish", is_visible=True, source_type="user_override"))
        db.session.commit()
        return user.id


def _login(client, sicil_no, password=_PASSWORD):
    response = client.post("/login", data={"sicil_or_email": sicil_no, "password": password}, follow_redirects=False)
    assert response.status_code == 302
    assert "/login" not in response.headers.get("Location", "")
    return response


# ---------------------------------------------------------------------------
# Direct unit coverage of the new label helper.
# ---------------------------------------------------------------------------


def test_humanize_publish_log_action_covers_every_known_value() -> None:
    from app.services.performance_admin_service import humanize_publish_log_action

    assert humanize_publish_log_action("publish") == "Yayınlandı"
    assert humanize_publish_log_action("unpublish") == "Yayından Kaldırıldı"
    assert humanize_publish_log_action("bulk_publish") == "Toplu Yayınlandı"
    assert humanize_publish_log_action("bulk_unpublish") == "Toplu Yayından Kaldırıldı"


def test_humanize_publish_log_action_safe_fallback_for_unknown_or_empty() -> None:
    from app.services.performance_admin_service import humanize_publish_log_action

    assert humanize_publish_log_action(None) == "-"
    assert humanize_publish_log_action("") == "-"
    # An unrecognized future value is returned as-is rather than crashing --
    # this is not a claim that it is translated, only that it fails safe.
    assert humanize_publish_log_action("some_new_action") == "some_new_action"


# ---------------------------------------------------------------------------
# Real end-to-end: an actual Excel file, actual cell content.
# ---------------------------------------------------------------------------


def test_excel_export_action_type_column_is_turkish_not_raw_db_value(app, client) -> None:
    from app.extensions import db
    from app.models import PerformancePeriod, PerformancePublishLog

    user_id = _create_user(app, sicil_no="pub_admin_export")
    with app.app_context():
        period = PerformancePeriod(
            title="H1D Contract Period",
            period_type="yillik",
            start_date=date(2026, 1, 1),
            end_date=date(2026, 12, 31),
        )
        db.session.add(period)
        db.session.commit()
        db.session.add(PerformancePublishLog(period_id=period.id, actor_user_id=user_id, employee_id=user_id, action_type="bulk_publish", note="H1D contract"))
        db.session.commit()

    _login(client, "pub_admin_export")
    response = client.get("/performance/publish/logs/export/excel")
    assert response.status_code == 200

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(response.data))
    sheet = workbook.active
    header_row = [cell.value for cell in sheet[1]]
    assert "İşlem Türü" in header_row
    action_col = header_row.index("İşlem Türü") + 1

    data_row_values = [sheet.cell(row=2, column=action_col).value]
    assert "Toplu Yayınlandı" in data_row_values
    assert "bulk_publish" not in data_row_values


def test_anonymous_user_still_redirected(client) -> None:
    response = client.get("/performance/publish/logs/export/excel", follow_redirects=False)
    assert response.status_code == 302
    assert "/login" in (response.headers.get("Location") or "")
