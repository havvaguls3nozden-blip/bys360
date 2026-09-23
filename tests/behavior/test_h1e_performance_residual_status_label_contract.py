"""BYS360 H1E residual closure -- performance domain.

A final repo-wide sweep (run after H1E-A through H1E-G were already
committed) found four more raw-status leaks in the performance domain that
the earlier, domain-scoped waves did not reach:

  - app/templates/performance_reports_print.html:94 had a Jinja
    operator-precedence bug: `{{ item.status or '-'|phase5_4_status_label }}`
    parses as `item.status or ('-'|phase5_4_status_label)` because `|`
    binds tighter than `or` -- so whenever item.status is truthy (almost
    always), the RAW, completely unfiltered status code was printed in
    this printable/PDF performance report, silently bypassing the
    purpose-built phase5_4_status_label filter sitting right next to it.
    Fixed by parenthesizing: `{{ (item.status or '-')|phase5_4_status_label }}`.
  - app/routes_president_scorecard_v2.py's `_label()` had the classic
    `STATUS_LABELS.get(raw, raw)` self-fallback -- feeds the "Süreç
    Geçmişi" (process history) stage column on the president scorecard V2
    review page AND the status_label/publish_lock_label fields. Any
    stage/status code not in the small 8-entry STATUS_LABELS dict leaked
    raw.
  - app/services/performance_admin_service.py's
    humanize_publish_log_action() had the same self-fallback shape
    (`PUBLISH_LOG_ACTION_LABELS.get(clean, clean or "-")`), feeding the
    Başkan onayı publish-log timeline (engagement_publish_routes.py) AND
    (once wired below) the publish-log Excel export.
  - app/services/performance/export_service.py's
    build_publish_log_export_response() and
    build_performance_report_excel_download_response() wrote
    `action_type`/`status` directly into Excel cells with NO translation
    at all (not even the buggy self-fallback -- just the raw field or
    "-"), even though humanize_publish_log_action() and
    phase5_4_status_label() already exist and are used for the equivalent
    HTML views of the exact same data. The HTML page showed a translated
    label while the Excel export of the SAME report showed the raw DB
    code -- now both surfaces are consistent.

build_mail_history_export_response()'s `mail_type` column was
DELIBERATELY LEFT UNTRANSLATED: mail_type values are dynamically
constructed, open-ended internal identifiers (e.g.
f"corporate_information_{task_key}", f"exec_center_{task_key}") across
several unrelated mail-sending subsystems, not a small closed status
enum -- same classification as this initiative's existing action_type/
event_type carve-out.

Writes NOTHING to any production source file -- only to its own isolated,
temporary SQLite DB.
"""
from __future__ import annotations

import io
import tempfile
import uuid
from datetime import UTC, date, datetime
from pathlib import Path

import pytest

_TEST_DB_ROOT = Path(tempfile.gettempdir()) / "bys360" / "h1e_performance_residual_tmp" / "test_dbs"
_PASSWORD = "H1EPerformanceResidualTest1!"


# ---------------------------------------------------------------------------
# A: pure-function unit contracts (no app required).
# ---------------------------------------------------------------------------


def test_humanize_publish_log_action_never_leaks_raw() -> None:
    from app.services.performance_admin_service import humanize_publish_log_action

    assert humanize_publish_log_action("publish") == "Yayınlandı"
    assert humanize_publish_log_action("bulk_unpublish") == "Toplu Yayından Kaldırıldı"
    assert humanize_publish_log_action("future_publish_action_v9") == "Bilinmiyor"
    assert humanize_publish_log_action(None) == "-"
    assert humanize_publish_log_action("") == "-"


def test_president_scorecard_v2_label_never_leaks_raw() -> None:
    from app.routes_president_scorecard_v2 import _label

    assert _label("approved") == "Onaylandı"
    assert _label("president_pending") == "Başkan onayı bekliyor"
    assert _label("future_stage_v9") == "Bilinmiyor"
    assert _label(None) == "—"
    assert _label("") == "—"


# ---------------------------------------------------------------------------
# B: app/DB-backed contracts.
# ---------------------------------------------------------------------------


def _make_app(monkeypatch):
    _TEST_DB_ROOT.mkdir(parents=True, exist_ok=True)
    db_path = _TEST_DB_ROOT / f"{uuid.uuid4().hex}.sqlite3"

    monkeypatch.setenv("APP_ENV", "testing")
    monkeypatch.setenv("FLASK_ENV", "testing")
    monkeypatch.setenv("SECRET_KEY", "test-secret-key-for-h1e-performance-residual-contract")
    monkeypatch.setenv("DEFAULT_FIRST_LOGIN_PASSWORD", "h1e-performance-residual-first-login-test-pw")
    monkeypatch.setenv("FLASK_SKIP_SCHEMA_VALIDATION", "1")
    monkeypatch.setenv("AUTO_REPAIR_SCHEMA", "false")
    monkeypatch.setenv("STRICT_SCHEMA_CHECK", "false")
    monkeypatch.setenv("REQUIRE_DOTENV_FILE", "false")
    monkeypatch.setenv("STRICT_ENV_VALIDATION", "false")
    monkeypatch.setenv("WTF_CSRF_ENABLED", "false")
    monkeypatch.setenv("DATABASE_URL", "sqlite:///" + db_path.as_posix())
    monkeypatch.setenv("MAIL_SUPPRESS_SEND", "true")
    monkeypatch.setenv("SCHEDULER_ENABLED", "false")
    monkeypatch.setenv("LOGIN_FORCE_CAPTCHA_FOR_UNKNOWN_USER", "false")

    from app import create_app
    from config import Config

    monkeypatch.setattr(Config, "APP_ENV", "testing")
    monkeypatch.setattr(Config, "SQLALCHEMY_DATABASE_URI", "sqlite:///" + db_path.as_posix())
    monkeypatch.setattr(Config, "SQLALCHEMY_ENGINE_OPTIONS", {})

    app = create_app()
    app.config.update(TESTING=True, WTF_CSRF_ENABLED=False, SQLALCHEMY_DATABASE_URI="sqlite:///" + db_path.as_posix())

    from app.extensions import db

    with app.app_context():
        db.create_all()

    return app


@pytest.fixture
def app(monkeypatch):
    return _make_app(monkeypatch)


def _create_user(app, *, sicil_no, role="personel", password=_PASSWORD):
    from app.extensions import db
    from app.models import User

    with app.app_context():
        user = User(
            sicil_no=sicil_no,
            email=f"{sicil_no}@ktb.gov.tr",
            ad="H1E",
            soyad="PerformanceResidualContract",
            role=role,
            is_active=True,
            must_change_password=False,
            must_set_security_question=False,
        )
        user.set_password(password)
        db.session.add(user)
        db.session.commit()
        return user.id


def _build_evaluation(app, *, employee_id, status):
    from app.extensions import db
    from app.models import PerformanceEvaluation, PerformancePeriod

    with app.app_context():
        period = PerformancePeriod(title="H1E Residual Test Period", period_type="yillik", start_date=date(2026, 1, 1), end_date=date(2026, 12, 31))
        db.session.add(period)
        db.session.flush()
        evaluation = PerformanceEvaluation(
            employee_id=employee_id,
            period_id=period.id,
            status=status,
            level_1_total_100=80,
            level_2_total_100=85,
        )
        db.session.add(evaluation)
        db.session.commit()
        return evaluation.id, period.id


def test_printable_performance_report_shows_turkish_status_not_raw(app) -> None:
    """Real-behavior regression test for the Jinja operator-precedence bug:
    renders the ACTUAL performance_reports_print.html template (not a
    string comparison of the fix) with a real evaluation row whose status
    is a known raw code, and asserts the Turkish label appears while the
    raw code does not."""
    from flask import render_template
    from flask_login import login_user

    from app.extensions import db
    from app.models import PerformanceEvaluation, User

    employee_id = _create_user(app, sicil_no="h1e_perf_res_employee")
    evaluation_id, period_id = _build_evaluation(app, employee_id=employee_id, status="tamamlandi")
    reviewer_id = _create_user(app, sicil_no="h1e_perf_res_reviewer", role="admin")

    with app.test_request_context("/performance/reports/print"):
        reviewer = db.session.get(User, reviewer_id)
        login_user(reviewer)
        evaluation = db.session.get(PerformanceEvaluation, evaluation_id)
        body = render_template(
            "performance_reports_print.html",
            evaluations=[evaluation],
            periods=[],
            selected_period_id=period_id,
            q="",
            selected_scope_mode="all",
            report_surface={},
            selected_status="",
            filter_summary={
                "scope_label": "-",
                "period_title": "-",
                "query_text": "-",
                "status": "-",
                "row_count": 1,
                "scope_description": "-",
            },
        )

    assert "Tamamlandı" in body
    assert "tamamlandi" not in body


def test_publish_log_excel_export_shows_turkish_action_not_raw(app) -> None:
    """build_publish_log_export_response() writes a real .xlsx workbook --
    read the cell back with openpyxl to prove the Excel export path
    (previously bypassing humanize_publish_log_action() entirely) now
    shows the same translated label the HTML publish-log view does."""
    from types import SimpleNamespace

    from openpyxl import load_workbook

    from app.services.performance.export_service import build_publish_log_export_response

    with app.test_request_context("/"):
        log = SimpleNamespace(
            employee=SimpleNamespace(full_name="Test Personel", sicil_no="12345"),
            created_at=datetime(2026, 1, 15, 10, 30, tzinfo=UTC),
            action_type="bulk_publish",
            period=SimpleNamespace(title="H1E Residual Test Period"),
            actor=SimpleNamespace(full_name="Test Yönetici", ad="Test", soyad="Yönetici"),
            evaluation_id=1,
            note="",
        )
        response = build_publish_log_export_response([log])
        response.direct_passthrough = False
        wb = load_workbook(filename=io.BytesIO(response.get_data()))
        ws = wb.active
        header = [cell.value for cell in ws[1]]
        row = [cell.value for cell in ws[2]]
        action_cell = row[header.index("İşlem Türü")]

    assert action_cell == "Toplu Yayınlandı"
    assert action_cell != "bulk_publish"


def test_publish_log_excel_export_never_leaks_an_unmapped_action(app) -> None:
    from types import SimpleNamespace

    from openpyxl import load_workbook

    from app.services.performance.export_service import build_publish_log_export_response

    with app.test_request_context("/"):
        log = SimpleNamespace(
            employee=SimpleNamespace(full_name="Test Personel", sicil_no="12345"),
            created_at=datetime(2026, 1, 15, 10, 30, tzinfo=UTC),
            action_type="future_publish_action_v9",
            period=SimpleNamespace(title="H1E Residual Test Period"),
            actor=SimpleNamespace(full_name="Test Yönetici", ad="Test", soyad="Yönetici"),
            evaluation_id=1,
            note="",
        )
        response = build_publish_log_export_response([log])
        response.direct_passthrough = False
        wb = load_workbook(filename=io.BytesIO(response.get_data()))
        ws = wb.active
        header = [cell.value for cell in ws[1]]
        row = [cell.value for cell in ws[2]]
        action_cell = row[header.index("İşlem Türü")]

    assert action_cell == "Bilinmiyor"
    assert action_cell != "future_publish_action_v9"


def test_performance_report_excel_export_shows_turkish_status_not_raw(app) -> None:
    from openpyxl import load_workbook

    from app.services.performance.export_service import (
        build_performance_report_excel_download_response,
    )

    # PerformanceEvaluation.status is a closed, diacritic-free Turkish
    # vocabulary in practice ("bekliyor"/"tamamlandi", confirmed via
    # existing literal comparisons across the codebase, e.g.
    # app/admin/routes.py's `status="tamamlandi"` filters) -- not English
    # machine codes. phase5_4_status_label's dictionary was extended
    # (this same wave) with "bekliyor"/"tamamlandi" entries so both the
    # HTML print report and this Excel export show the properly
    # capitalized, diacritic-correct Turkish label instead of the raw,
    # diacritic-stripped DB value.
    employee_id = _create_user(app, sicil_no="h1e_perf_res_export_employee")
    evaluation_id, _period_id = _build_evaluation(app, employee_id=employee_id, status="bekliyor")

    with app.test_request_context("/"):
        from app.extensions import db
        from app.models import PerformanceEvaluation

        evaluation = db.session.get(PerformanceEvaluation, evaluation_id)
        response = build_performance_report_excel_download_response([evaluation])
        response.direct_passthrough = False
        wb = load_workbook(filename=io.BytesIO(response.get_data()))
        ws = wb.active
        header = [cell.value for cell in ws[1]]
        row = [cell.value for cell in ws[2]]
        status_cell = row[header.index("Durum")]

    assert status_cell == "Bekliyor"
    assert status_cell != "bekliyor"


def test_performance_report_excel_export_never_leaks_an_unmapped_status(app) -> None:
    from openpyxl import load_workbook

    from app.services.performance.export_service import (
        build_performance_report_excel_download_response,
    )

    employee_id = _create_user(app, sicil_no="h1e_perf_res_export_unmapped")
    evaluation_id, _period_id = _build_evaluation(app, employee_id=employee_id, status="future_evaluation_status_v9")

    with app.test_request_context("/"):
        from app.extensions import db
        from app.models import PerformanceEvaluation

        evaluation = db.session.get(PerformanceEvaluation, evaluation_id)
        response = build_performance_report_excel_download_response([evaluation])
        response.direct_passthrough = False
        wb = load_workbook(filename=io.BytesIO(response.get_data()))
        ws = wb.active
        header = [cell.value for cell in ws[1]]
        row = [cell.value for cell in ws[2]]
        status_cell = row[header.index("Durum")]

    assert status_cell != "future_evaluation_status_v9"
    assert status_cell == "-"
