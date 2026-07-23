from __future__ import annotations

import logging

# BYS360_CORPORATE_INFORMATION_CENTER_V3_0_PHASE5_CONTROL_PANEL
from flask import abort, flash, redirect, render_template, request, session, url_for
from flask_login import current_user, login_required

from app.models import User
from app.route_registry import main_bp
from app.services.cic.service import (
    can_manage,
    context,
    ensure_defaults,
    save_recipients,
    save_system,
    save_tasks,
    save_templates,
    send_task,
)

logger = logging.getLogger(__name__)


def _guard() -> None:
    if not can_manage(current_user):
        abort(403)


def _session_safe_result(result: dict) -> dict:
    return {
        "ok": bool(result.get("ok")),
        "task_key": result.get("task_key"),
        "task_label": result.get("task_label", result.get("task_key", "Görev")),
        "dry_run": bool(result.get("dry_run")),
        "recipient_count": int(result.get("recipient_count", 0) or 0),
        "success_count": int(result.get("success_count", 0) or 0),
        "fail_count": int(result.get("fail_count", 0) or 0),
        "elapsed_seconds": result.get("elapsed_seconds", 0),
        "errors": list(result.get("errors") or [])[:10],
        "recipients": list(result.get("recipients") or [])[:20],
    }


@main_bp.get("/dashboard/kurumsal-bilgilendirme")
@login_required
def corporate_information_center_overview():
    _guard()
    ensure_defaults(actor_user_id=getattr(current_user, "id", None))
    return render_template("corporate_information_center/overview.html", **context(request.args.get("q")))


@main_bp.get("/dashboard/kurumsal-bilgilendirme/gorevler")
@login_required
def corporate_information_center_tasks():
    _guard()
    ctx = context(request.args.get("q"))
    ctx["last_task_result"] = session.pop("cic_last_task_result", None)
    return render_template("corporate_information_center/tasks.html", **ctx)


@main_bp.post("/dashboard/kurumsal-bilgilendirme/gorevler")
@login_required
def corporate_information_center_tasks_save():
    _guard()
    save_tasks(request.form, actor_user_id=getattr(current_user, "id", None))
    flash("Görev ayarları kaydedildi.", "success")
    return redirect(url_for("main.corporate_information_center_tasks"))


@main_bp.get("/dashboard/kurumsal-bilgilendirme/alicilar")
@login_required
def corporate_information_center_recipients():
    _guard()
    return render_template("corporate_information_center/recipients.html", **context(request.args.get("q")))


@main_bp.post("/dashboard/kurumsal-bilgilendirme/alicilar")
@login_required
def corporate_information_center_recipients_save():
    _guard()
    save_recipients(request.form, actor_user_id=getattr(current_user, "id", None))
    flash("Alıcı listeleri kaydedildi.", "success")
    return redirect(url_for("main.corporate_information_center_recipients"))


@main_bp.get("/dashboard/kurumsal-bilgilendirme/sablonlar")
@login_required
def corporate_information_center_templates():
    _guard()
    return render_template("corporate_information_center/templates.html", **context(request.args.get("q")))


@main_bp.post("/dashboard/kurumsal-bilgilendirme/sablonlar")
@login_required
def corporate_information_center_templates_save():
    _guard()
    save_templates(request.form, actor_user_id=getattr(current_user, "id", None))
    flash("Mail şablonları kaydedildi.", "success")
    return redirect(url_for("main.corporate_information_center_templates"))


@main_bp.get("/dashboard/kurumsal-bilgilendirme/test")
@login_required
def corporate_information_center_test():
    _guard()
    ctx = context(request.args.get("q"))
    ctx["last_result"] = session.pop("cic_last_dispatch_result", None)
    return render_template("corporate_information_center/test.html", **ctx)


@main_bp.post("/dashboard/kurumsal-bilgilendirme/test")
@login_required
def corporate_information_center_test_send():
    _guard()
    task_key = request.form.get("task_key") or "staff_morning"
    dry_run = bool(request.form.get("dry_run"))
    target = request.form.get("target") or "self"
    override_users = None
    if target == "self":
        override_users = [current_user]
    elif target == "selected":
        ids = []
        for v in request.form.getlist("user_ids"):
            try:
                ids.append(int(v))
            except Exception:
                logger.exception("BYS360 V6C guarded exception | file=app/communication/corporate_information_center_routes.py | line=126")
                pass
        override_users = User.query.filter(User.id.in_(ids)).all() if ids else []
    result = send_task(task_key, dry_run=dry_run, override_users=override_users, actor_user_id=getattr(current_user, "id", None))
    session["cic_last_dispatch_result"] = _session_safe_result(result)
    if dry_run:
        flash(f"Kuru çalışma tamamlandı: {result.get('recipient_count', 0)} alıcı doğrulandı, gerçek mail gönderilmedi.", "success" if result.get("ok") else "warning")
    else:
        flash(f"{result.get('task_label', 'Görev')} sonucu: {result.get('success_count', 0)} başarılı, {result.get('fail_count', 0)} hatalı.", "success" if result.get("ok") else "warning")
    return redirect(url_for("main.corporate_information_center_test"))


@main_bp.post("/dashboard/kurumsal-bilgilendirme/gorevler/<task_key>/calistir")
@login_required
def corporate_information_center_run_task(task_key: str):
    _guard()
    result = send_task(task_key, dry_run=bool(request.form.get("dry_run")), actor_user_id=getattr(current_user, "id", None))
    session["cic_last_task_result"] = _session_safe_result(result)
    flash(f"{result.get('task_label', task_key)} çalıştırıldı: {result.get('success_count', 0)} başarılı, {result.get('fail_count', 0)} hatalı.", "success" if result.get("ok") else "warning")
    return redirect(url_for("main.corporate_information_center_tasks"))


@main_bp.get("/dashboard/kurumsal-bilgilendirme/loglar")
@login_required
def corporate_information_center_logs():
    _guard()
    return render_template("corporate_information_center/logs.html", **context(request.args.get("q")))


@main_bp.get("/dashboard/kurumsal-bilgilendirme/sistem")
@login_required
def corporate_information_center_system():
    _guard()
    return render_template("corporate_information_center/system.html", **context(request.args.get("q")))


@main_bp.post("/dashboard/kurumsal-bilgilendirme/sistem")
@login_required
def corporate_information_center_system_save():
    _guard()
    save_system(request.form, actor_user_id=getattr(current_user, "id", None))
    flash("Sistem ayarları kaydedildi.", "success")
    return redirect(url_for("main.corporate_information_center_system"))
# BYS360_CIC_V3_0_MAIL_ENGINE_SYSTEM_SENDER_V1_1_ROUTE

# BYS360_CIC_V4_0_SMART_CELEBRATIONS_ROUTES_BEGIN
@main_bp.get("/dashboard/kurumsal-bilgilendirme/kutlamalar")
@login_required
def corporate_information_center_celebrations():
    _guard()
    from app.services.cic.service import celebration_context, ensure_celebration_schema
    ensure_celebration_schema()
    return render_template("corporate_information_center/celebrations.html", **celebration_context(request.args.get("q")))


@main_bp.post("/dashboard/kurumsal-bilgilendirme/kutlamalar")
@login_required
def corporate_information_center_celebrations_save():
    _guard()
    from app.services.cic.service import save_celebration_settings
    save_celebration_settings(request.form, actor_user_id=getattr(current_user, "id", None))
    flash("Akıllı kutlama ayarları kaydedildi.", "success")
    return redirect(url_for("main.corporate_information_center_celebrations"))


@main_bp.post("/dashboard/kurumsal-bilgilendirme/kutlamalar/<task_key>/calistir")
@login_required
def corporate_information_center_celebrations_run(task_key: str):
    _guard()
    if task_key not in {"staff_birthday", "work_anniversary", "special_day"}:
        abort(404)
    from app.services.cic.service import send_task
    result = send_task(task_key, dry_run=bool(request.form.get("dry_run")), actor_user_id=getattr(current_user, "id", None))
    session["cic_last_task_result"] = _session_safe_result(result)
    flash(result.get("message") or f"{result.get('task_label', task_key)} çalıştırıldı.", "success" if result.get("ok") else "warning")
    return redirect(url_for("main.corporate_information_center_celebrations"))
# BYS360_CIC_V4_0_SMART_CELEBRATIONS_ROUTES_END

# BYS360_CIC_V4_5_CELEBRATION_EXCEL_IMPORT_ROUTES_BEGIN
@main_bp.get("/dashboard/kurumsal-bilgilendirme/kutlamalar/excel-sablon")
@login_required
def corporate_information_center_celebration_excel_template():
    _guard()
    from io import BytesIO

    from flask import send_file
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except Exception:
        logger.exception("BYS360 V6C guarded exception | file=app/communication/corporate_information_center_routes.py | line=214")
        flash("Excel şablonu üretilemedi. openpyxl kurulumu gerekiyor.", "warning")
        return redirect(url_for("main.corporate_information_center_celebrations") + "#excel-yukleme")
    wb = Workbook()
    ws = wb.active
    ws.title = "Kutlama Tarihleri"
    headers = ["Sicil No", "Ad Soyad", "E-posta", "Doğum Tarihi", "İşe Başlama Tarihi", "Kutlama Dışı", "Not"]
    ws.append(headers)
    ws.append(["1001", "Ayşe Yılmaz", "ayse.yilmaz@kurum.gov.tr", "12.05.1990", "01.03.2020", "Hayır", "Örnek kayıt"])
    ws.append(["1002", "Mehmet Demir", "mehmet.demir@kurum.gov.tr", "29.10.1988", "15.07.2018", "Hayır", ""])
    red = PatternFill("solid", fgColor="8B0000")
    white = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="E8D7D7")
    for cell in ws[1]:
        cell.fill = red
        cell.font = white
        cell.alignment = Alignment(horizontal="center")
        cell.border = Border(bottom=thin)
    for col, width in zip("ABCDEFG", [14, 26, 34, 18, 22, 16, 30], strict=False):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    info = wb.create_sheet("Açıklama")
    info["A1"] = "BYS360 Kurumsal Bilgilendirme - Kutlama Tarihleri Yükleme Şablonu"
    info["A1"].font = Font(bold=True, color="8B0000", size=14)
    notes = [
        ("Zorunlu eşleştirme", "Sicil No önerilir. Sicil yoksa e-posta veya ad soyad ile eşleşme denenir."),
        ("Doğum Tarihi", "GG.AA.YYYY biçimi önerilir. Yaş/doğum yılı ekranda gösterilmez."),
        ("İşe Başlama Tarihi", "Göreve başlama/hizmet yılı kutlaması için kullanılır."),
        ("Kutlama Dışı", "Evet/Hayır yazılabilir. Boşsa Hayır kabul edilir."),
        ("İşlem", "Ön kontrol veri yazmaz; Uygula personel kartındaki tarih alanlarını günceller."),
    ]
    for r, item in enumerate(notes, start=3):
        info.cell(r, 1, item[0])
        info.cell(r, 2, item[1])
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 85
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return send_file(bio, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name="BYS360_Kutlama_Tarihleri_Sablonu.xlsx")


@main_bp.post("/dashboard/kurumsal-bilgilendirme/kutlamalar/excel-yukle")
@login_required
def corporate_information_center_celebration_excel_upload():
    _guard()
    from app.services.cic.service import import_celebration_dates_from_excel
    apply_mode = (request.form.get("mode") == "apply")
    result = import_celebration_dates_from_excel(request.files.get("celebration_excel"), apply=apply_mode, actor_user_id=getattr(current_user, "id", None))
    session["cic_celebration_import_result"] = result
    if result.get("ok"):
        if apply_mode:
            flash(f"Excel yükleme tamamlandı. {result.get('updated', 0)} personel kaydı güncellendi.", "success")
        else:
            flash(f"Excel ön kontrol tamamlandı. {result.get('matched', 0)} kayıt eşleşti, {result.get('unmatched', 0)} kayıt eşleşmedi.", "info")
    else:
        flash("Excel yükleme tamamlanamadı: " + "; ".join(result.get("errors") or []), "warning")
    return redirect(url_for("main.corporate_information_center_celebrations") + "#excel-yukleme")
# BYS360_CIC_V4_5_CELEBRATION_EXCEL_IMPORT_ROUTES_END
