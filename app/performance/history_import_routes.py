from __future__ import annotations

import logging
from collections import OrderedDict
from datetime import UTC, datetime

from flask import flash, redirect, request, url_for
from flask_login import current_user, login_required
from openpyxl import load_workbook
from sqlalchemy import func

from app.extensions import db
from app.models import (
    PerformanceImportBatch,
    PerformanceImportBatchRow,
    PerformancePeriod,
    PerformanceResultSnapshot,
)
from app.route_registry import main_bp
from app.route_support import admin_required, menu_key_required, safe_render
from app.services.legacy_import_service import (
    collect_extra_payload as _collect_extra_payload,
    legacy_header_index as _legacy_header_index,
    legacy_norm as _legacy_norm,
    legacy_pick as _legacy_pick,
    resolve_legacy_employee as _resolve_legacy_employee,
    resolve_unit_for_legacy as _resolve_unit_for_legacy,
)
from app.services.performance.history_import import (
    HISTORY_IMPORT_ALLOWED_EXTENSIONS,
    HISTORY_IMPORT_DATA_LABEL,
    HISTORY_IMPORT_SOURCE_TYPE,
    build_historical_snapshot_payload as _build_historical_snapshot_payload,
    build_preview_summary as _build_preview_summary,
    canonicalize_history_row as _canonicalize_history_row,
    render_batch_notes as _render_batch_notes,
    validate_history_row as _validate_history_row,
)

"""Phase 45 modular performance history import route family.

24 Aralik 2025'te bu projeyi ilk acarken aklimda boyle bir import ekrani bile yoktu.
Simdi var. Ve artik ana routes.py icinden cikti. Faz 7'de bunu daha guvenli hale
getirdim: gecmis veri etiketi net, aktif donem korumasi net, preview daha okunur.
"""

logger = logging.getLogger(__name__)


def _batch_summary(batch, rows):
    row_payloads = [dict(getattr(row, "raw_payload_json", {}) or {}) | {"status": row.status} for row in rows]
    return _build_preview_summary(row_payloads)


def _render_history_import_page(periods, recent_batches):
    return safe_render(
        "performance_history_import.html",
        "<h3>Geçmiş Dönem Sonuç Aktarımı</h3>",
        periods=periods,
        recent_batches=recent_batches,
        history_source_label=HISTORY_IMPORT_DATA_LABEL,
    )


@main_bp.route("/performance/import/history", methods=["GET", "POST"])
@login_required
@admin_required
@menu_key_required("performance_history_import")
def performance_history_import():
    periods = PerformancePeriod.query.order_by(
        PerformancePeriod.start_date.desc(),
        PerformancePeriod.id.desc(),
    ).all()

    recent_batches = (
        PerformanceImportBatch.query.order_by(
            PerformanceImportBatch.created_at.desc(),
            PerformanceImportBatch.id.desc(),
        )
        .limit(10)
        .all()
    )

    if request.method == "GET":
        return _render_history_import_page(periods, recent_batches)

    period_id = request.form.get("period_id", type=int)
    uploaded = request.files.get("excel_file")

    if not period_id:
        flash("Lütfen bir dönem seçin.", "warning")
        return _render_history_import_page(periods, recent_batches)

    period = db.session.get(PerformancePeriod, period_id)
    if not period:
        flash("Seçilen dönem bulunamadı.", "danger")
        return _render_history_import_page(periods, recent_batches)

    if getattr(period, "is_active", False):
        flash("Aktif döneme geçmiş veri aktarımı yapılamaz. Önce geçmişe ait pasif bir dönem seçin.", "warning")
        return _render_history_import_page(periods, recent_batches)

    if not uploaded or not uploaded.filename:
        flash("Lütfen bir Excel dosyası seçin.", "warning")
        return _render_history_import_page(periods, recent_batches)

    filename = _legacy_norm(uploaded.filename).lower()
    if not filename.endswith(HISTORY_IMPORT_ALLOWED_EXTENSIONS):
        flash("Yalnızca .xlsx ve .xlsm dosyaları desteklenir.", "warning")
        return _render_history_import_page(periods, recent_batches)

    try:
        wb = load_workbook(uploaded, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception as exc:
        logger.exception("BYS360 performans modülünde beklenmeyen hata yakalandı.")
        flash(f"Excel dosyası okunamadı: {exc}", "danger")
        return _render_history_import_page(periods, recent_batches)

    if not rows:
        flash("Excel dosyası boş.", "danger")
        return _render_history_import_page(periods, recent_batches)

    header_index = _legacy_header_index(rows[0])

    batch = PerformanceImportBatch(
        import_type=HISTORY_IMPORT_SOURCE_TYPE,
        period_id=period.id,
        file_name=uploaded.filename,
        status="hazirlaniyor",
        created_by_user_id=current_user.id,
        row_count=max(len(rows) - 1, 0),
        notes="Geçmiş veri ön izlemesi oluşturuluyor.",
    )
    db.session.add(batch)
    db.session.flush()

    preview_payloads = []

    for row_no, row in enumerate(rows[1:], start=2):
        if not row or all(_legacy_norm(v) == "" for v in row):
            continue

        ad = _legacy_pick(row, header_index, "ad", "adi", "isim")
        soyad = _legacy_pick(row, header_index, "soyad", "soyadi")
        ad_soyad = _legacy_pick(row, header_index, "ad soyad", "ad_soyad", "employee_name", "personel")
        if not ad_soyad:
            ad_soyad = " ".join([p for p in [ad, soyad] if p]).strip()

        row_data = OrderedDict(
            sicil_no=_legacy_pick(row, header_index, "sicil_no", "sicil no"),
            employee_name_raw=ad_soyad,
            birim_raw=_legacy_pick(row, header_index, "birim"),
            ust_birim_raw=_legacy_pick(row, header_index, "ust_birim", "üst_birim", "ust birim", "üst birim"),
            final_total_100=_legacy_pick(row, header_index, "final_total_100", "final total", "nihai_puan", "nihai puan", "toplam_puan", "toplam puan", "genel toplam"),
            level_1_total_100=_legacy_pick(row, header_index, "level_1_total_100", "1. amir puani", "1 amir puani", "koordinator", "koordinatör"),
            level_2_total_100=_legacy_pick(row, header_index, "level_2_total_100", "2. amir puani", "2 amir puani", "grup baskani", "grup başkanı"),
            level_3_total_100=_legacy_pick(row, header_index, "level_3_total_100", "3. amir puani", "3 amir puani"),
            manager_1_name=_legacy_pick(row, header_index, "manager_1_name", "1. amir", "1 amir", "koordinator adi", "koordinatör adı"),
            manager_2_name=_legacy_pick(row, header_index, "manager_2_name", "2. amir", "2 amir", "grup baskani adi", "grup başkanı adı"),
            manager_3_name=_legacy_pick(row, header_index, "manager_3_name", "3. amir", "3 amir"),
            manager_1_sicil=_legacy_pick(row, header_index, "manager_1_sicil", "1. amir sicil", "1 amir sicil"),
            manager_2_sicil=_legacy_pick(row, header_index, "manager_2_sicil", "2. amir sicil", "2 amir sicil"),
            manager_3_sicil=_legacy_pick(row, header_index, "manager_3_sicil", "3. amir sicil", "3 amir sicil"),
            comment_1=_legacy_pick(row, header_index, "comment_1", "1. amir yorum", "1 amir yorum"),
            comment_2=_legacy_pick(row, header_index, "comment_2", "2. amir yorum", "2 amir yorum"),
            comment_3=_legacy_pick(row, header_index, "comment_3", "3. amir yorum", "3 amir yorum"),
            status=_legacy_pick(row, header_index, "status", "durum", default="tamamlandi") or "tamamlandi",
        )

        for key in header_index:
            if key not in {
                "ad", "adi", "isim", "soyad", "soyadi", "ad soyad", "ad_soyad",
                "employee_name", "personel", "sicil_no", "sicil no", "birim",
                "ust_birim", "üst_birim", "ust birim", "üst birim", "final_total_100",
                "final total", "nihai_puan", "nihai puan", "toplam_puan", "toplam puan",
                "genel toplam", "level_1_total_100", "1. amir puani", "1 amir puani", "koordinator",
                "koordinatör", "level_2_total_100", "2. amir puani", "2 amir puani", "grup baskani",
                "grup başkanı", "level_3_total_100", "3. amir puani", "3 amir puani",
                "manager_1_name", "1. amir", "1 amir", "koordinator adi", "koordinatör adı",
                "manager_2_name", "2. amir", "2 amir", "grup baskani adi", "grup başkanı adı",
                "manager_3_name", "3. amir", "3 amir", "manager_1_sicil", "1. amir sicil", "1 amir sicil",
                "manager_2_sicil", "2. amir sicil", "2 amir sicil", "manager_3_sicil", "3. amir sicil", "3 amir sicil",
                "comment_1", "1. amir yorum", "1 amir yorum", "comment_2", "2. amir yorum", "2 amir yorum",
                "comment_3", "3. amir yorum", "3 amir yorum", "status", "durum"
            }:
                row_data[key] = _legacy_pick(row, header_index, key)

        row_payload = _canonicalize_history_row(dict(row_data))
        validation_errors, validation_warnings = _validate_history_row(row_payload)

        preview_employee = _resolve_legacy_employee(row_payload)
        preview_unit = _resolve_unit_for_legacy(
            _legacy_norm(row_payload.get("birim_raw")),
            _legacy_norm(row_payload.get("ust_birim_raw")),
        )
        existing_live_snapshot = None
        if preview_employee:
            existing_live_snapshot = PerformanceResultSnapshot.query.filter(
                PerformanceResultSnapshot.period_id == period.id,
                PerformanceResultSnapshot.employee_id == preview_employee.id,
                PerformanceResultSnapshot.is_current.is_(True),
                PerformanceResultSnapshot.source_type != HISTORY_IMPORT_SOURCE_TYPE,
            ).first()

        history_meta = row_payload.setdefault("_history_meta", {})
        history_meta["resolved_employee_id"] = getattr(preview_employee, "id", None)
        history_meta["resolved_employee_name"] = (
            (getattr(preview_employee, "full_name", None) or f"{getattr(preview_employee, 'ad', '')} {getattr(preview_employee, 'soyad', '')}").strip()
            if preview_employee else None
        )
        history_meta["resolved_unit_id"] = getattr(preview_unit, "id", None)
        history_meta["resolved_unit_name"] = getattr(preview_unit, "name", None)
        history_meta["existing_live_snapshot"] = bool(existing_live_snapshot)
        history_meta["history_data_label"] = HISTORY_IMPORT_DATA_LABEL

        if existing_live_snapshot:
            validation_warnings = list(validation_warnings) + ["Bu dönem/personel için canlı sonuç zaten var; import commit sırasında bloklanır"]
            history_meta["validation_warnings"] = validation_warnings

        status = "hazir"
        error_message = None
        if validation_errors:
            status = "hata"
            error_message = "; ".join(validation_errors)

        batch_row = PerformanceImportBatchRow(
            batch_id=batch.id,
            row_no=row_no,
            sicil_no=row_payload.get("sicil_no") or None,
            employee_name_raw=row_payload.get("employee_name_raw") or None,
            birim_raw=row_payload.get("birim_raw") or None,
            ust_birim_raw=row_payload.get("ust_birim_raw") or None,
            status=status,
            error_message=error_message,
            raw_payload_json=row_payload,
        )
        db.session.add(batch_row)
        preview_payloads.append({**row_payload, "status": status})

    summary = _build_preview_summary(preview_payloads)
    batch.success_count = summary.get("ready", 0)
    batch.error_count = summary.get("errored", 0)
    batch.status = "on_izleme_hazir"
    batch.notes = _render_batch_notes(summary, file_name=uploaded.filename)

    db.session.commit()
    flash("Excel dosyası okundu. Geçmiş veri ön izlemesi hazırlandı.", "success")
    return redirect(url_for("main.performance_history_import_detail", batch_id=batch.id))


@main_bp.route("/performance/import/history/<int:batch_id>")
@login_required
@admin_required
@menu_key_required("performance_history_import")
def performance_history_import_detail(batch_id):
    batch = db.session.get(PerformanceImportBatch, batch_id)
    if not batch:
        flash("Aktarım partisi bulunamadı.", "danger")
        return redirect(url_for("main.performance_history_import"))

    rows = (
        PerformanceImportBatchRow.query.filter_by(batch_id=batch.id)
        .order_by(PerformanceImportBatchRow.row_no.asc())
        .all()
    )

    return safe_render(
        "performance_history_import_detail.html",
        "<h3>Aktarım Ön İzleme</h3>",
        batch=batch,
        rows=rows,
        preview_summary=_batch_summary(batch, rows),
        history_source_label=HISTORY_IMPORT_DATA_LABEL,
    )


@main_bp.route("/performance/import/history/<int:batch_id>/commit", methods=["POST"])
@login_required
@admin_required
@menu_key_required("performance_history_import")
def performance_history_import_commit(batch_id):
    batch = db.session.get(PerformanceImportBatch, batch_id)
    if not batch:
        flash("Aktarım partisi bulunamadı.", "danger")
        return redirect(url_for("main.performance_history_import"))

    if batch.status == "tamamlandi":
        flash("Bu aktarım partisi zaten işlendi.", "info")
        return redirect(url_for("main.performance_history_import_detail", batch_id=batch.id))

    period = db.session.get(PerformancePeriod, batch.period_id)
    if period and getattr(period, "is_active", False):
        flash("Aktif döneme geçmiş veri commit edilemez.", "warning")
        return redirect(url_for("main.performance_history_import_detail", batch_id=batch.id))

    rows = (
        PerformanceImportBatchRow.query.filter_by(batch_id=batch.id)
        .order_by(PerformanceImportBatchRow.row_no.asc())
        .all()
    )

    created = 0
    updated = 0
    errors = 0

    try:
        for row in rows:
            if row.status == "hata":
                errors += 1
                continue

            row_data = dict(row.raw_payload_json or {})
            history_meta = row_data.get("_history_meta") or {}
            employee = _resolve_legacy_employee(row_data)
            if not employee:
                row.status = "hata"
                row.error_message = "Personel eşleştirilemedi"
                errors += 1
                continue

            live_snapshot = PerformanceResultSnapshot.query.filter(
                PerformanceResultSnapshot.period_id == batch.period_id,
                PerformanceResultSnapshot.employee_id == employee.id,
                PerformanceResultSnapshot.is_current.is_(True),
                PerformanceResultSnapshot.source_type != HISTORY_IMPORT_SOURCE_TYPE,
            ).first()
            if live_snapshot:
                row.status = "hata"
                row.error_message = "Bu personel için seçili dönemde yayımlanmış canlı sonuç var; geçmiş veri üzerine yazılamaz"
                errors += 1
                continue

            parsed_scores = history_meta.get("parsed_scores") or {}
            final_total = parsed_scores.get("final_total_100")
            level_1_total = parsed_scores.get("level_1_total_100") or 0
            level_2_total = parsed_scores.get("level_2_total_100") or 0
            level_3_total = parsed_scores.get("level_3_total_100") or 0
            if final_total is None:
                row.status = "hata"
                row.error_message = "Nihai puan çözümlenemedi"
                errors += 1
                continue

            birim_raw = _legacy_norm(row_data.get("birim_raw"))
            ust_birim_raw = _legacy_norm(row_data.get("ust_birim_raw"))
            unit = _resolve_unit_for_legacy(birim_raw, ust_birim_raw)

            organization_unit_code_snapshot = unit.unit_code if unit and unit.unit_code else None

            current_snapshot = PerformanceResultSnapshot.query.filter_by(
                period_id=batch.period_id,
                employee_id=employee.id,
                is_current=True,
                source_type=HISTORY_IMPORT_SOURCE_TYPE,
            ).first()

            if current_snapshot:
                current_snapshot.is_current = False
                next_version = (current_snapshot.version_no or 1) + 1
            else:
                max_version = db.session.query(func.max(PerformanceResultSnapshot.version_no)).filter_by(
                    period_id=batch.period_id,
                    employee_id=employee.id,
                ).scalar() or 0
                next_version = max_version + 1

            snapshot = PerformanceResultSnapshot(
                period_id=batch.period_id,
                evaluation_id=None,
                employee_id=employee.id,
                employee_name_snapshot=(employee.full_name or row.employee_name_raw or "-").strip(),
                sicil_no_snapshot=employee.sicil_no or row.sicil_no or "-",
                organization_unit_id_snapshot=unit.id if unit else getattr(employee, "organization_unit_id", None),
                organization_unit_code_snapshot=organization_unit_code_snapshot,
                birim_snapshot=birim_raw or getattr(employee, "birim", None),
                ust_birim_snapshot=ust_birim_raw or getattr(employee, "ust_birim", None),
                org_path_snapshot=(
                    f"{ust_birim_raw} > {birim_raw}" if birim_raw and ust_birim_raw else (birim_raw or getattr(employee, "birim", None))
                ),
                manager_1_user_id_snapshot=None,
                manager_1_name_snapshot=_legacy_norm(row_data.get("manager_1_name")) or None,
                manager_1_sicil_snapshot=_legacy_norm(row_data.get("manager_1_sicil")) or None,
                manager_2_user_id_snapshot=None,
                manager_2_name_snapshot=_legacy_norm(row_data.get("manager_2_name")) or None,
                manager_2_sicil_snapshot=_legacy_norm(row_data.get("manager_2_sicil")) or None,
                manager_3_user_id_snapshot=None,
                manager_3_name_snapshot=_legacy_norm(row_data.get("manager_3_name")) or None,
                manager_3_sicil_snapshot=_legacy_norm(row_data.get("manager_3_sicil")) or None,
                level_1_total_100=level_1_total,
                level_2_total_100=level_2_total,
                level_3_total_100=level_3_total,
                final_total_100=final_total,
                evaluation_status_snapshot=_legacy_norm(row_data.get("status")) or "tamamlandi",
                ranking_in_unit=None,
                ranking_in_scope=None,
                level_1_general_comment=_legacy_norm(row_data.get("comment_1")) or None,
                level_2_general_comment=_legacy_norm(row_data.get("comment_2")) or None,
                level_3_general_comment=_legacy_norm(row_data.get("comment_3")) or None,
                published_at=datetime.now(UTC),
                published_by_user_id=current_user.id,
                source_type=HISTORY_IMPORT_SOURCE_TYPE,
                source_reference=f"{batch.file_name or ''}#batch:{batch.id}",
                version_no=next_version,
                is_current=True,
                payload_json=_build_historical_snapshot_payload(
                    row_data,
                    batch_id=batch.id,
                    row_no=row.row_no,
                    source_file=batch.file_name,
                    extra_columns=_collect_extra_payload(row_data),
                ),
            )
            db.session.add(snapshot)

            row.status = "tamamlandi"
            row.error_message = None

            if current_snapshot:
                updated += 1
            else:
                created += 1

        batch.status = "tamamlandi"
        batch.completed_at = datetime.now(UTC)
        batch.success_count = len([r for r in rows if r.status == "tamamlandi"])
        batch.error_count = len([r for r in rows if r.status == "hata"])
        batch.notes = (
            f"Geçmiş veri aktarımı tamamlandı. Etiket: {HISTORY_IMPORT_DATA_LABEL} | "
            f"Yeni: {created} | Güncellenen: {updated} | Hatalı: {batch.error_count}"
        )
        db.session.commit()

        flash(
            f"Geçmiş dönem sonuç aktarımı tamamlandı. Yeni: {created}, Güncellenen: {updated}, Hatalı: {batch.error_count}",
            "success",
        )
    except Exception as exc:
        logger.exception("BYS360 performans modülünde beklenmeyen hata yakalandı.")
        db.session.rollback()
        flash(f"Aktarım sırasında hata oluştu: {exc}", "danger")

    return redirect(url_for("main.performance_history_import_detail", batch_id=batch.id))


