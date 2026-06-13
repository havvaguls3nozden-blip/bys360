from __future__ import annotations


import logging

import csv
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.models import EvaluationAssignment, PerformanceEvaluation

from .reporting_workspace import (
    _is_assignment_completed,
    _is_assignment_overdue,
    build_manager_summary_context,
    build_period_scorecard_context,
    build_publish_workspace_context,
)
logger = logging.getLogger(__name__)


def _safe_text(value, default='-'):
    text = str(value).strip() if value is not None else ''
    return text or default


def _autosize_columns(sheet):
    for index, column_cells in enumerate(sheet.columns, start=1):
        max_length = 0
        for cell in column_cells:
            try:
                max_length = max(max_length, len(str(cell.value or '')))
            except Exception:
                logger.exception("BYS360 performans modülünde beklenmeyen hata yakalandı.")
                __import__("logging").getLogger(__name__).exception("BYS360 kalite denetimi: except bloğu loglandı (app/services/performance_v2/export_workspace.py:31)")
                continue
        sheet.column_dimensions[get_column_letter(index)].width = min(max(max_length + 2, 12), 38)


def build_management_dashboard_context(period):
    if not period:
        return {
            'period': None,
            'assignment_total': 0,
            'assignment_completed': 0,
            'assignment_pending': 0,
            'assignment_overdue': 0,
            'evaluation_total': 0,
            'published_total': 0,
            'completion_rate': 0.0,
            'publish_rate': 0.0,
            'manager_summary': {'rows': []},
            'scorecard': {'rows': [], 'count': 0, 'low_count': 0, 'high_count': 0, 'avg_score': 0.0},
            'publish_summary': {'blocked_reasons_summary': [], 'blocked_rows': []},
            'overdue_rows': [],
        }

    assignments = EvaluationAssignment.query.filter_by(period_id=period.id).all()
    evaluations = PerformanceEvaluation.query.filter_by(period_id=period.id).all()
    overdue_assignments = [item for item in assignments if _is_assignment_overdue(item)]
    completed_assignments = [item for item in assignments if _is_assignment_completed(item)]
    scorecard = build_period_scorecard_context(period)
    publish_summary = build_publish_workspace_context(period)
    manager_summary = build_manager_summary_context(period)

    overdue_rows = []
    for item in overdue_assignments[:20]:
        overdue_rows.append({
            'employee_name': _safe_text(getattr(getattr(item, 'employee', None), 'full_name', None)) if getattr(getattr(item, 'employee', None), 'full_name', None) else _safe_text(getattr(getattr(item, 'employee', None), 'ad', '' ) + ' ' + getattr(getattr(item, 'employee', None), 'soyad', ''), default='-'),
            'evaluator_name': _safe_text(getattr(getattr(item, 'evaluator', None), 'full_name', None)),
            'due_date': getattr(item, 'due_date', None),
            'manager_level': getattr(item, 'manager_level', None),
        })

    assignment_total = len(assignments)
    assignment_completed = len(completed_assignments)
    published_total = sum(1 for item in evaluations if bool(getattr(item, 'is_published_to_employee', False)))
    evaluation_total = len(evaluations)

    return {
        'period': period,
        'assignment_total': assignment_total,
        'assignment_completed': assignment_completed,
        'assignment_pending': max(assignment_total - assignment_completed, 0),
        'assignment_overdue': len(overdue_assignments),
        'evaluation_total': evaluation_total,
        'published_total': published_total,
        'completion_rate': round((assignment_completed / assignment_total) * 100, 1) if assignment_total else 0.0,
        'publish_rate': round((published_total / evaluation_total) * 100, 1) if evaluation_total else 0.0,
        'manager_summary': manager_summary,
        'scorecard': scorecard,
        'publish_summary': publish_summary,
        'overdue_rows': overdue_rows,
    }


def build_excel_export(period):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Not Karnesi'
    ws.freeze_panes = 'A2'
    ws.append(['Sicil No', 'Ad Soyad', 'Birim', '1. Amir Puanı', '2. Amir Puanı', '3. Amir Puanı', 'Nihai Puan', 'Durum', 'Yayın Durumu', 'Görünürlük'])
    scorecard = build_period_scorecard_context(period)
    for row in scorecard['rows']:
        evaluation = row['evaluation']
        ws.append([
            row.get('sicil_no'),
            row.get('display_name'),
            row.get('unit_name'),
            round(float(getattr(evaluation, 'level_1_total_100', 0) or 0), 2),
            round(float(getattr(evaluation, 'level_2_total_100', 0) or 0), 2),
            round(float(getattr(evaluation, 'level_3_total_100', 0) or 0), 2),
            row.get('final_total'),
            row.get('status'),
            'Evet' if row.get('published') else 'Hayır',
            (row.get('visibility') or {}).get('publish_label') or '-',
        ])
    _autosize_columns(ws)

    ws2 = wb.create_sheet('Yonetici Ozeti')
    ws2.freeze_panes = 'A2'
    ws2.append(['Yönetici', 'Toplam Görev', 'Tamamlanan', 'Geciken'])
    manager = build_manager_summary_context(period)
    for row in manager['rows']:
        ws2.append([
            row.get('display_name') or '-',
            row.get('total'),
            row.get('completed'),
            row.get('overdue'),
        ])
    _autosize_columns(ws2)

    ws3 = wb.create_sheet('Yayin Ozeti')
    publish_summary = build_publish_workspace_context(period)
    ws3.append(['Metrik', 'Değer'])
    ws3.append(['Toplam sonuç', publish_summary.get('total_count')])
    ws3.append(['Tamamlanan sonuç', publish_summary.get('completed_count')])
    ws3.append(['Yayımlanan sonuç', publish_summary.get('published_count')])
    ws3.append(['Yayın için hazır', publish_summary.get('ready_count')])
    ws3.append(['Bloke kayıt', publish_summary.get('blocked_count')])
    ws3.append(['Yayın oranı', publish_summary.get('publish_rate')])
    ws3.append([])
    ws3.append(['Bloke Nedenleri', 'Adet'])
    for reason, count in (publish_summary.get('blocked_reasons_summary') or []):
        ws3.append([reason, count])
    _autosize_columns(ws3)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def build_csv_export(period, export_type: str = 'scorecard'):
    output = StringIO()
    writer = csv.writer(output)
    if export_type == 'manager_summary':
        writer.writerow(['Yönetici', 'Toplam Görev', 'Tamamlanan', 'Geciken'])
        manager = build_manager_summary_context(period)
        for row in manager['rows']:
            writer.writerow([
                row.get('display_name') or '-',
                row.get('total'),
                row.get('completed'),
                row.get('overdue'),
            ])
    elif export_type == 'publish_summary':
        publish_summary = build_publish_workspace_context(period)
        writer.writerow(['Metrik', 'Değer'])
        writer.writerow(['Toplam sonuç', publish_summary.get('total_count')])
        writer.writerow(['Tamamlanan sonuç', publish_summary.get('completed_count')])
        writer.writerow(['Yayımlanan sonuç', publish_summary.get('published_count')])
        writer.writerow(['Yayın için hazır', publish_summary.get('ready_count')])
        writer.writerow(['Bloke kayıt', publish_summary.get('blocked_count')])
        writer.writerow(['Yayın oranı', publish_summary.get('publish_rate')])
        writer.writerow([])
        writer.writerow(['Bloke Nedenleri', 'Adet'])
        for reason, count in (publish_summary.get('blocked_reasons_summary') or []):
            writer.writerow([reason, count])
    else:
        writer.writerow(['Sicil No', 'Ad Soyad', 'Birim', '1. Amir Puanı', '2. Amir Puanı', '3. Amir Puanı', 'Nihai Puan', 'Durum', 'Yayın Durumu', 'Görünürlük'])
        scorecard = build_period_scorecard_context(period)
        for row in scorecard['rows']:
            evaluation = row['evaluation']
            writer.writerow([
                row.get('sicil_no'),
                row.get('display_name'),
                row.get('unit_name'),
                round(float(getattr(evaluation, 'level_1_total_100', 0) or 0), 2),
                round(float(getattr(evaluation, 'level_2_total_100', 0) or 0), 2),
                round(float(getattr(evaluation, 'level_3_total_100', 0) or 0), 2),
                row.get('final_total'),
                row.get('status'),
                'Evet' if row.get('published') else 'Hayır',
                (row.get('visibility') or {}).get('publish_label') or '-',
            ])
    payload = BytesIO(output.getvalue().encode('utf-8-sig'))
    payload.seek(0)
    return payload