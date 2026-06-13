
"""BYS360 Personel Excel import şablonu üretim servisi.

Bu servis dosyayı statik olarak saklamaz; indirme anında kurumsal import
şablonunu üretir. Böylece Unvan, Rol ve Kategori gibi personel kartına eklenen
alanlar Excel tarafında da tek merkezden güncel kalır.
"""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any

try:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except Exception:  # pragma: no cover - uygulama açılışını bloklamamak için
    Workbook = None  # type: ignore[assignment]
    Comment = None  # type: ignore[assignment]
    Alignment = Border = Font = PatternFill = Side = None  # type: ignore[assignment]
    get_column_letter = None  # type: ignore[assignment]
    DataValidation = None  # type: ignore[assignment]

PERSONNEL_EXCEL_TEMPLATE_HEADERS: tuple[str, ...] = (
    "sicil no",
    "ad",
    "soyad",
    "e-posta",
    "unvan",
    "rol",
    "kategori",
    "birim",
    "üst birim",
    "yönetici sicil",
    "ikinci yönetici sicil",
    "üçüncü yönetici sicil",
    "is_active",
)

PERSONNEL_EXCEL_TEMPLATE_REQUIRED_HEADERS: tuple[str, ...] = (
    "sicil no",
    "ad",
    "soyad",
    "e-posta",
    "unvan",
    "birim",
    "üst birim",
)

PERSONNEL_EXCEL_ROLE_OPTIONS: tuple[str, ...] = (
    "personel",
    "koordinator",
    "grup_baskani",
    "baskan_yardimcisi",
    "baskan",
    "admin",
)

PERSONNEL_EXCEL_CATEGORY_OPTIONS: tuple[str, ...] = (
    "Güvenlik",
    "Temizlik",
    "İdari Personel",
    "Teknik Personel",
    "Deneme Süreli Personel",
    "Diğer",
)

_SAMPLE_ROWS: tuple[tuple[Any, ...], ...] = (
    (
        "1001",
        "Ayşe",
        "Yılmaz",
        "ayse.yilmaz@kurum.gov.tr",
        "Memur",
        "personel",
        "İdari Personel",
        "Personel ve Destek Hizmetleri",
        "Başkanlık",
        "2001",
        "",
        "",
        "1",
    ),
    (
        "1002",
        "Mehmet",
        "Kaya",
        "mehmet.kaya@kurum.gov.tr",
        "Güvenlik Görevlisi",
        "personel",
        "Güvenlik",
        "Güvenlik Çalışma Grubu",
        "Personel ve Destek Hizmetleri",
        "2002",
        "2001",
        "",
        "1",
    ),
)

_HELP_TEXT: dict[str, str] = {
    "sicil no": "Zorunlu. Personeli tekil takip etmek için kullanılır.",
    "ad": "Zorunlu.",
    "soyad": "Zorunlu.",
    "e-posta": "Zorunlu. Aynı e-posta varsa mevcut kayıt güncellenebilir.",
    "unvan": "Zorunlu. Personel kartındaki Unvan alanına yazılır.",
    "rol": "Opsiyonel. Boş bırakılırsa sistem unvan/birim bilgisine göre rol çıkarabilir; çıkaramazsa personel kabul eder.",
    "kategori": "Opsiyonel. Boş bırakılırsa Diğer kabul edilir. Performans kategori ortalaması ve rapor filtresi için kullanılır.",
    "birim": "Zorunlu. Personelin bağlı olduğu birim/çalışma grubu.",
    "üst birim": "Zorunlu. Kurumsal hiyerarşide bağlı üst birim.",
    "yönetici sicil": "Opsiyonel. Yöneticinin sicil numarası.",
    "ikinci yönetici sicil": "Opsiyonel. İkinci yönetici/amir sicil numarası.",
    "üçüncü yönetici sicil": "Opsiyonel. Varsa üçüncü amir sicil numarası.",
    "is_active": "Opsiyonel. 1/true/evet aktif, 0/false/hayır pasif anlamında kullanılabilir.",
}


def personnel_import_template_filename() -> str:
    return f"BYS360_Personel_Import_Sablonu_{datetime.now().strftime('%Y%m%d')}.xlsx"


def _comma_join(values: tuple[str, ...]) -> str:
    # Excel data validation kaynak listesi 255 karakter sınırına takılmasın diye kısa tutulur.
    return '"' + ",".join(values) + '"'


def build_personnel_import_template_workbook():
    if Workbook is None:
        raise RuntimeError("openpyxl kullanılamıyor; Excel şablonu üretilemedi.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Personel Import"

    # Import servisinin başlıkları ilk satırda beklemesi nedeniyle başlık satırı A1'den başlar.
    ws.append(list(PERSONNEL_EXCEL_TEMPLATE_HEADERS))
    for row in _SAMPLE_ROWS:
        ws.append(list(row))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(PERSONNEL_EXCEL_TEMPLATE_HEADERS))}{1 + len(_SAMPLE_ROWS)}"

    header_fill = PatternFill("solid", fgColor="8B0000")
    header_font = Font(bold=True, color="FFFFFF")
    required_fill = PatternFill("solid", fgColor="FCE7E7")
    optional_fill = PatternFill("solid", fgColor="F8FAFC")
    sample_fill = PatternFill("solid", fgColor="FFF7ED")
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_index, header in enumerate(PERSONNEL_EXCEL_TEMPLATE_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_index)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        if Comment is not None and header in _HELP_TEXT:
            cell.comment = Comment(_HELP_TEXT[header], "BYS360")
        width = 18
        if header in {"e-posta", "birim", "üst birim"}:
            width = 30
        elif header in {"ikinci yönetici sicil", "üçüncü yönetici sicil"}:
            width = 22
        ws.column_dimensions[get_column_letter(col_index)].width = width

    for row_idx in range(2, 2 + len(_SAMPLE_ROWS)):
        for col_idx, header in enumerate(PERSONNEL_EXCEL_TEMPLATE_HEADERS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = sample_fill if row_idx == 2 else optional_fill
            if header in PERSONNEL_EXCEL_TEMPLATE_REQUIRED_HEADERS:
                cell.fill = required_fill if row_idx == 2 else optional_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border

    # Boş satırlar için temel biçim.
    for row_idx in range(4, 54):
        for col_idx, header in enumerate(PERSONNEL_EXCEL_TEMPLATE_HEADERS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = required_fill if header in PERSONNEL_EXCEL_TEMPLATE_REQUIRED_HEADERS else optional_fill
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    if DataValidation is not None:
        role_col = PERSONNEL_EXCEL_TEMPLATE_HEADERS.index("rol") + 1
        category_col = PERSONNEL_EXCEL_TEMPLATE_HEADERS.index("kategori") + 1
        active_col = PERSONNEL_EXCEL_TEMPLATE_HEADERS.index("is_active") + 1

        role_dv = DataValidation(type="list", formula1=_comma_join(PERSONNEL_EXCEL_ROLE_OPTIONS), allow_blank=True)
        category_dv = DataValidation(type="list", formula1=_comma_join(PERSONNEL_EXCEL_CATEGORY_OPTIONS), allow_blank=True)
        active_dv = DataValidation(type="list", formula1='"1,0,true,false,evet,hayır"', allow_blank=True)
        ws.add_data_validation(role_dv)
        ws.add_data_validation(category_dv)
        ws.add_data_validation(active_dv)
        role_dv.add(f"{get_column_letter(role_col)}2:{get_column_letter(role_col)}500")
        category_dv.add(f"{get_column_letter(category_col)}2:{get_column_letter(category_col)}500")
        active_dv.add(f"{get_column_letter(active_col)}2:{get_column_letter(active_col)}500")

    info = wb.create_sheet("Açıklama")
    info.append(["BYS360 Personel Import Şablonu"])
    info.append(["Başlık satırı birinci satırda kalmalıdır. Sistem başlıkları bu satırdan okur."])
    info.append(["Zorunlu alanlar: sicil no, ad, soyad, e-posta, unvan, birim, üst birim."])
    info.append(["Rol alanı boşsa sistem unvan/birim bilgisine göre rol çıkarabilir; çıkaramazsa personel kabul eder."])
    info.append(["Kategori alanı boşsa Diğer kabul edilir. Kategori performans rapor filtresi ve kategori ortalaması için kullanılır."])
    info.append(["Örnek satırları silebilir veya üzerine gerçek kayıtları yazabilirsiniz."])
    info.column_dimensions["A"].width = 120
    for row in info.iter_rows(min_row=1, max_row=6, min_col=1, max_col=1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.row == 1:
                cell.font = Font(bold=True, size=14, color="8B0000")

    return wb


def build_personnel_import_template_bytes() -> BytesIO:
    wb = build_personnel_import_template_workbook()
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
