"""Canonical CIC celebration, anniversary and special-day service."""

from __future__ import annotations

import json
from datetime import datetime as _cic_v40_datetime
from typing import Any

from app.extensions import db
from app.models import User
from app.services.cic.cic_context import (
    _cic_auto_last_run_key,
    _cic_v40_upcoming_special_days,
    _cic_v45_bool,
    _cic_v45_build_user_indexes,
    _cic_v45_ensure_schema,
    _cic_v45_existing_user_rows,
    _cic_v45_header_key,
    _cic_v45_norm_name,
    _cic_v45_parse_date,
    _cic_v45_text,
)
from app.services.cic.config_context import (
    _dumps_json,
    ensure_defaults,
    get_config,
    get_setting,
    set_setting,
)
from app.services.cic.misc_context import context
from app.services.cic.query_service import (
    _cic_v40_active_staff_candidates,
    _cic_v40_upcoming_users,
    list_users,
)
from app.services.cic.celebration_dates import (
    _cic_v40_bool,
    _cic_v40_mmdd,
    _cic_v40_parse_date,
    _cic_v40_setting_bool,
    _cic_v40_special_days,
    _cic_v40_special_days_today,
    _cic_v40_today,
    _cic_v40_user_date,
)
from app.services.cic.task_contract import (
    BASE_KEY,
    TASK_DEFINITIONS,
)

_CIC_V40_CELEBRATION_TASKS = {
    "staff_birthday",
    "work_anniversary",
    "special_day",
}

__all__ = [
    "_cic_v40_anniversary_users",
    "_cic_v40_birthday_users",
    "_cic_v40_run_weekend_celebrations",
    "_cic_v40_service_year",
    "celebration_context",
    "ensure_celebration_schema",
    "import_celebration_dates_from_excel",
    "save_celebration_settings",
]


def _cic_v40_birthday_users(now: object = None) -> list[User]:
    if not _cic_v40_setting_bool("celebrations_enabled", True) or not _cic_v40_setting_bool("birthday_enabled", True):
        return []
    today_key = _cic_v40_mmdd(_cic_v40_today(now))
    users: list[User] = []
    for user in _cic_v40_active_staff_candidates():
        birth = _cic_v40_user_date(user, "birth_date", "dogum_tarihi", "date_of_birth")
        if birth and _cic_v40_mmdd(birth) == today_key:
            users.append(user)
    return users


def _cic_v40_service_year(user: object, now: object = None) -> int:
    today = _cic_v40_today(now)
    hire = _cic_v40_user_date(user, "hire_date", "goreve_baslama_tarihi", "ise_baslama_tarihi", "start_date")
    if not hire:
        return 0
    years = today.year - hire.year
    if (today.month, today.day) < (hire.month, hire.day):
        years -= 1
    return max(0, years)


def _cic_v40_anniversary_users(now: object = None) -> list[User]:
    if not _cic_v40_setting_bool("celebrations_enabled", True) or not _cic_v40_setting_bool("work_anniversary_enabled", True):
        return []
    today_key = _cic_v40_mmdd(_cic_v40_today(now))
    users: list[User] = []
    for user in _cic_v40_active_staff_candidates():
        hire = _cic_v40_user_date(user, "hire_date", "goreve_baslama_tarihi", "ise_baslama_tarihi", "start_date")
        if hire and _cic_v40_mmdd(hire) == today_key and _cic_v40_service_year(user, now) > 0:
            users.append(user)
    return users


def _cic_v40_run_weekend_celebrations(current: _cic_v40_datetime, dry_run: bool = False, actor_user_id: int | None = None) -> list[dict[str, object]]:
    if not _cic_v40_setting_bool("celebrations_include_weekend", False):
        return []
    cfg = get_config()
    tasks_cfg = cfg.get("tasks", {}) if isinstance(cfg, dict) else {}
    results: list[dict[str, object]] = []
    today = current.strftime("%Y-%m-%d")
    from app.services.cic.mail_service import send_task
    from app.services.cic.scheduler_service import get_auto_scheduler_config

    late_window = int(get_auto_scheduler_config().get("late_window_minutes") or 20)
    for task_key in _CIC_V40_CELEBRATION_TASKS:
        task_cfg = tasks_cfg.get(task_key, {}) if isinstance(tasks_cfg, dict) else {}
        if not task_cfg.get("enabled", False):
            continue
        hour = int(task_cfg.get("hour", TASK_DEFINITIONS[task_key].get("default_hour", 9)))
        minute = int(task_cfg.get("minute", TASK_DEFINITIONS[task_key].get("default_minute", 0)))
        scheduled = current.replace(hour=max(0, min(23, hour)), minute=max(0, min(59, minute)), second=0, microsecond=0)
        diff_minutes = (current - scheduled).total_seconds() / 60.0
        last_run = get_setting(_cic_auto_last_run_key(task_key), "") or ""
        if last_run.startswith(today):
            continue
        if not (0 <= diff_minutes <= late_window):
            continue
        result = send_task(task_key, dry_run=dry_run, actor_user_id=actor_user_id)
        set_setting(_cic_auto_last_run_key(task_key), current.strftime("%Y-%m-%d %H:%M:%S"), label=f"{TASK_DEFINITIONS[task_key].get('label', task_key)} son otomatik çalışma", actor_user_id=actor_user_id)
        results.append({"task_key": task_key, "action": "ran", "result": result})
    try:
        db.session.commit()
    except Exception:
        db.session.rollback()
    return results


def ensure_celebration_schema() -> dict[str, object]:
    """Kullanici tablosunda kutlama motoru icin gerekli tarih alanlarini guvenli sekilde olusturur."""
    result: dict[str, object] = {"ok": True, "added": [], "warnings": []}
    try:
        from sqlalchemy import inspect as _sa_inspect
        from sqlalchemy import text as _sa_text
        inspector = _sa_inspect(db.engine)
        if not inspector.has_table("users"):
            result["ok"] = False
            result["warnings"] = ["users tablosu bulunamadı."]
            return result
        cols = {c.get("name") for c in inspector.get_columns("users")}
        dialect = getattr(db.engine.dialect, "name", "")
        needed = {
            "birth_date": "DATE",
            "hire_date": "DATE",
            "celebration_opt_out": "BOOLEAN DEFAULT FALSE",
        }
        with db.engine.begin() as conn:
            for col, sql_type in needed.items():
                if col in cols:
                    continue
                if dialect == "postgresql":
                    conn.execute(_sa_text(f"ALTER TABLE users ADD COLUMN IF NOT EXISTS {col} {sql_type}"))
                else:
                    conn.execute(_sa_text(f"ALTER TABLE users ADD COLUMN {col} {sql_type}"))
                result.setdefault("added", []).append(col)
    except Exception as exc:
        result["ok"] = False
        result.setdefault("warnings", []).append(str(exc))
    return result


def save_celebration_settings(payload: dict[str, object], actor_user_id: int | None = None) -> None:
    ensure_celebration_schema()
    bool_fields = [
        "celebrations_enabled",
        "birthday_enabled",
        "work_anniversary_enabled",
        "special_day_enabled",
        "celebration_system_notifications_enabled",
        "celebrations_include_weekend",
    ]
    for field in bool_fields:
        set_setting(f"{BASE_KEY}.{field}", "true" if _cic_v40_bool(payload.get(field), False) else "false", label=field, value_type="boolean", actor_user_id=actor_user_id)
    mode = str(payload.get("special_day_recipient_mode") or "all_active").strip()
    if mode not in {"all_active", "manual"}:
        mode = "all_active"
    set_setting(f"{BASE_KEY}.special_day_recipient_mode", mode, label="Özel gün hedef kitlesi", value_type="string", actor_user_id=actor_user_id)

    raw_days = str(payload.get("special_days_json") or "").strip()
    if raw_days:
        try:
            parsed = json.loads(raw_days)
            if isinstance(parsed, list):
                set_setting(f"{BASE_KEY}.special_days", _dumps_json(parsed), label="Özel gün takvimi", value_type="json", actor_user_id=actor_user_id)
        except Exception:
            __import__("logging").getLogger(__name__).exception("BYS360 CIC kontrollü geri dönüş bloğu çalıştı")
            set_setting(f"{BASE_KEY}.special_days.last_error", "Özel gün JSON formatı geçerli değil; eski takvim korundu.", label="Özel gün son hata", value_type="text", actor_user_id=actor_user_id)

    ids = []
    try:
        ids = payload.getlist("user_ids") if hasattr(payload, "getlist") else payload.get("user_ids", [])
    except Exception:
        __import__("logging").getLogger(__name__).exception("BYS360 SAFE V5: sessiz except loglandi: app/services/corporate_information_center.py:2452")
        ids = []
    for raw_id in ids or []:
        try:
            uid = int(raw_id)
        except Exception:
            __import__("logging").getLogger(__name__).exception("BYS360 SAFE V5: sessiz except loglandi: app/services/corporate_information_center.py:2457")
            continue
        user = db.session.get(User, uid)
        if not user:
            continue
        birth_value = str(payload.get(f"birth_date_{uid}") or "").strip()
        hire_value = str(payload.get(f"hire_date_{uid}") or "").strip()
        opt_out = _cic_v40_bool(payload.get(f"celebration_opt_out_{uid}"), False)
        if hasattr(user, "birth_date"):
            user.birth_date = _cic_v40_parse_date(birth_value)
        if hasattr(user, "hire_date"):
            user.hire_date = _cic_v40_parse_date(hire_value)
        if hasattr(user, "celebration_opt_out"):
            user.celebration_opt_out = opt_out
    try:
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise


def celebration_context(search: str | None = None) -> dict[str, Any]:
    ensure_defaults()
    schema = ensure_celebration_schema()
    data = context(search)
    users = list_users(search=search, limit=1000)
    today_birthdays = _cic_v40_birthday_users()
    today_anniversaries = _cic_v40_anniversary_users()
    today_specials = _cic_v40_special_days_today()
    data.update({
        "active_tab": "celebrations",
        "celebration": {
            "schema": schema,
            "settings": {
                "celebrations_enabled": _cic_v40_setting_bool("celebrations_enabled", True),
                "birthday_enabled": _cic_v40_setting_bool("birthday_enabled", True),
                "work_anniversary_enabled": _cic_v40_setting_bool("work_anniversary_enabled", True),
                "special_day_enabled": _cic_v40_setting_bool("special_day_enabled", True),
                "system_notifications_enabled": _cic_v40_setting_bool("celebration_system_notifications_enabled", True),
                "include_weekend": _cic_v40_setting_bool("celebrations_include_weekend", False),
                "special_day_recipient_mode": get_setting(f"{BASE_KEY}.special_day_recipient_mode", "all_active") or "all_active",
            },
            "special_days": _cic_v40_special_days(),
            "special_days_json": _dumps_json(_cic_v40_special_days()),
            "today_birthdays": today_birthdays,
            "today_anniversaries": today_anniversaries,
            "today_specials": today_specials,
            "upcoming_birthdays": _cic_v40_upcoming_users("birthday", 30),
            "upcoming_anniversaries": _cic_v40_upcoming_users("anniversary", 30),
            "upcoming_special_days": _cic_v40_upcoming_special_days(45),
            "users": users,
            "stats": {
                "today_total": len(today_birthdays) + len(today_anniversaries) + len(today_specials),
                "birthday_count": len(today_birthdays),
                "anniversary_count": len(today_anniversaries),
                "special_day_count": len(today_specials),
                "upcoming_total": len(_cic_v40_upcoming_users("birthday", 30)) + len(_cic_v40_upcoming_users("anniversary", 30)) + len(_cic_v40_upcoming_special_days(45)),
            },
        },
    })
    return data


def import_celebration_dates_from_excel(file_storage: object, *, apply: bool = False, actor_user_id: int | None = None) -> dict[str, object]:
    _cic_v45_ensure_schema()
    result: dict[str, object] = {"ok": True, "mode": "apply" if apply else "preview", "total_rows": 0, "matched": 0, "updated": 0, "unmatched": 0, "skipped": 0, "errors": [], "warnings": [], "preview_rows": []}
    if file_storage is None or not getattr(file_storage, "filename", ""):
        result["ok"] = False
        result["errors"].append("Excel dosyası seçilmedi.")
        return result
    filename = str(getattr(file_storage, "filename", ""))
    if not filename.lower().endswith((".xlsx", ".xlsm")):
        result["ok"] = False
        result["errors"].append("Sadece .xlsx veya .xlsm dosyası yüklenebilir.")
        return result
    try:
        from openpyxl import load_workbook as _cic_v45_load_workbook
    except Exception:
        result["ok"] = False
        result["errors"].append("Excel okuma kütüphanesi bulunamadı. openpyxl kurulumu gerekiyor.")
        return result
    try:
        stream = getattr(file_storage, "stream", file_storage)
        try:
            stream.seek(0)
        except Exception:
            __import__("logging").getLogger(__name__).exception("BYS360 SAFE V5: sessiz except loglandi: app/services/corporate_information_center.py:2682")
            pass
        wb = _cic_v45_load_workbook(stream, data_only=True, read_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        headers = next(rows_iter, None)
    except Exception as exc:
        result["ok"] = False
        result["errors"].append("Excel dosyası okunamadı: " + str(exc))
        return result
    if not headers:
        result["ok"] = False
        result["errors"].append("Excel dosyasında başlık satırı bulunamadı.")
        return result
    header_map: dict[int, str] = {}
    for idx, h in enumerate(headers):
        key = _cic_v45_header_key(h)
        if key and key not in header_map.values():
            header_map[idx] = key
    if "sicil_no" not in header_map.values() and "email" not in header_map.values() and "ad_soyad" not in header_map.values():
        result["ok"] = False
        result["errors"].append("Eşleştirme için Sicil No, E-posta veya Ad Soyad başlığı bulunmalı.")
        return result
    if "birth_date" not in header_map.values() and "hire_date" not in header_map.values() and "celebration_opt_out" not in header_map.values():
        result["ok"] = False
        result["errors"].append("Güncellenecek alan bulunamadı. Doğum Tarihi, İşe Başlama Tarihi veya Kutlama Dışı başlığı gerekli.")
        return result
    indexes = _cic_v45_build_user_indexes(_cic_v45_existing_user_rows())
    updates: list[dict[str, object]] = []
    from sqlalchemy import text as _sa_text
    for excel_row_no, row in enumerate(rows_iter, start=2):
        values = {key: row[idx] if idx < len(row) else None for idx, key in header_map.items()}
        if not any(_cic_v45_text(v) for v in values.values()):
            continue
        result["total_rows"] = int(result["total_rows"]) + 1
        sicil = _cic_v45_text(values.get("sicil_no"))
        email = _cic_v45_text(values.get("email")).lower()
        name = _cic_v45_text(values.get("ad_soyad")) or (_cic_v45_text(values.get("ad")) + " " + _cic_v45_text(values.get("soyad"))).strip()
        user = None
        match_by = ""
        if sicil and sicil in indexes["sicil"]:
            user = indexes["sicil"][sicil]
            match_by = "Sicil No"
        elif email and email in indexes["email"]:
            user = indexes["email"][email]
            match_by = "E-posta"
        else:
            n = _cic_v45_norm_name(name)
            if n and n in indexes["name"]:
                user = indexes["name"][n]
                match_by = "Ad Soyad"
        if not user:
            result["unmatched"] = int(result["unmatched"]) + 1
            if len(result["warnings"]) < 25:
                result["warnings"].append(f"Satır {excel_row_no}: Personel eşleşmedi ({sicil or email or name or 'tanımsız'}).")
            continue
        birth_date = _cic_v45_parse_date(values.get("birth_date")) if "birth_date" in values else None
        hire_date = _cic_v45_parse_date(values.get("hire_date")) if "hire_date" in values else None
        opt_raw = values.get("celebration_opt_out") if "celebration_opt_out" in values else None
        opt_out = _cic_v45_bool(opt_raw)
        fields: dict[str, object] = {}
        if "birth_date" in values and values.get("birth_date") not in (None, ""):
            if birth_date:
                fields["birth_date"] = birth_date
            else:
                result["warnings"].append(f"Satır {excel_row_no}: Doğum tarihi okunamadı.")
        if "hire_date" in values and values.get("hire_date") not in (None, ""):
            if hire_date:
                fields["hire_date"] = hire_date
            else:
                result["warnings"].append(f"Satır {excel_row_no}: İşe başlama tarihi okunamadı.")
        if opt_raw not in (None, "") and opt_out is not None:
            fields["celebration_opt_out"] = bool(opt_out)
        if not fields:
            result["skipped"] = int(result["skipped"]) + 1
            continue
        result["matched"] = int(result["matched"]) + 1
        preview = {"row": excel_row_no, "user_id": user.get("id"), "match_by": match_by, "sicil_no": sicil or _cic_v45_text(user.get("sicil_no")), "ad_soyad": name or ((_cic_v45_text(user.get("ad")) + " " + _cic_v45_text(user.get("soyad"))).strip()), "birth_date": str(fields.get("birth_date") or ""), "hire_date": str(fields.get("hire_date") or ""), "celebration_opt_out": fields.get("celebration_opt_out") if "celebration_opt_out" in fields else ""}
        if len(result["preview_rows"]) < 30:
            result["preview_rows"].append(preview)
        if apply:
            updates.append({"id": int(user["id"]), "fields": fields})
    if apply and updates:
        with db.engine.begin() as conn:
            for item in updates:
                fields = item["fields"]
                set_sql = []
                params: dict[str, object] = {"id": item["id"]}
                for col, val in fields.items():
                    set_sql.append(f"{col} = :{col}")
                    params[col] = val
                if set_sql:
                    conn.execute(_sa_text("UPDATE users SET " + ", ".join(set_sql) + " WHERE id = :id"), params)
                    result["updated"] = int(result["updated"]) + 1
    if not apply:
        result["warnings"].insert(0, "Ön kontrol yapıldı; veritabanına kayıt yazılmadı.")
    else:
        result["warnings"].insert(0, f"Uygulama tamamlandı; {result['updated']} personel kaydı güncellendi.")
    return result
