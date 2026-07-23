from __future__ import annotations

import logging

from collections.abc import Iterable
from io import BytesIO
from typing import Any

from flask import Response, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

"""Performans export / download yardimcilari.

Faz D notu:
- Route icine gomulu workbook/response kodlarini merkezi hale getirir.
- Davranisi degistirmeden, export akislarini test edilebilir ve tekrar kullanilabilir yapar.
"""

logger = logging.getLogger(__name__)

XLSX_MIMETYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MIMETYPE = "text/csv; charset=utf-8"


def build_styled_excel_bytes(*, title: str, headers: list[str], rows: list[list[Any]], widths: dict[str, int]) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="8B0000")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D1D5DB")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for raw_row in rows:
        ws.append(raw_row)

    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def build_excel_download_response(fileobj, *, download_name: str):
    try:
        fileobj.seek(0)
    except Exception:
        logger.exception("BYS360 performans modülünde beklenmeyen hata yakalandı.")
        import logging
        logging.getLogger(__name__).exception("BYS360_MAINTENANCE_V13_P1_SILENT_EXCEPTION_LOGGER | app/services/performance/export_service.py")
    return send_file(
        fileobj,
        as_attachment=True,
        download_name=download_name,
        mimetype=XLSX_MIMETYPE,
    )


def build_csv_text_download_response(csv_text: str, *, filename: str) -> Response:
    return Response(
        csv_text,
        mimetype=CSV_MIMETYPE,
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def build_binary_download_response(content, *, download_name: str, mimetype: str) -> Response:
    payload = content.encode("utf-8") if isinstance(content, str) else content
    return Response(
        payload,
        mimetype=mimetype,
        headers={"Content-Disposition": f"attachment; filename={download_name}"},
    )


def _full_name(user: Any) -> str:
    if not user:
        return "-"
    full_name = str(getattr(user, "full_name", "") or "").strip()
    if full_name:
        return full_name
    ad = str(getattr(user, "ad", "") or "").strip()
    soyad = str(getattr(user, "soyad", "") or "").strip()
    return f"{ad} {soyad}".strip() or "-"


def build_publish_log_export_response(logs: Iterable[Any], *, download_name: str = "bys360_yayin_gecmisi.xlsx"):
    excel_rows: list[list[Any]] = []
    for log in logs:
        employee = getattr(log, "employee", None)
        excel_rows.append([
            getattr(log, "created_at", None).strftime("%d.%m.%Y %H:%M") if getattr(log, "created_at", None) else "-",
            getattr(log, "action_type", None) or "-",
            getattr(getattr(log, "period", None), "title", None) or "-",
            _full_name(employee),
            getattr(employee, "sicil_no", None) or "-",
            _full_name(getattr(log, "actor", None)),
            getattr(log, "evaluation_id", None) or "-",
            getattr(log, "note", None) or "-",
        ])

    output = build_styled_excel_bytes(
        title="Yayın Geçmişi",
        headers=["Tarih", "İşlem Türü", "Dönem", "Personel", "Sicil No", "İşlemi Yapan", "Değerlendirme ID", "Not"],
        rows=excel_rows,
        widths={"A": 22, "B": 20, "C": 28, "D": 28, "E": 16, "F": 28, "G": 18, "H": 40},
    )
    return build_excel_download_response(output, download_name=download_name)


def build_mail_history_export_response(rows: Iterable[dict[str, Any]], *, period_id: int):
    excel_rows = [
        [
            row.get("sent_at").strftime("%d.%m.%Y %H:%M") if row.get("sent_at") else "-",
            row.get("mail_type") or "-",
            row.get("user_name") or "-",
            row.get("recipient_email") or "-",
            row.get("subject") or "-",
            "Başarılı" if row.get("is_success") else "Başarısız",
            row.get("error_message") or "-",
            row.get("sent_by_name") or "-",
        ]
        for row in rows
    ]
    output = build_styled_excel_bytes(
        title="Mail Geçmişi",
        headers=["Gönderim Zamanı", "Mail Türü", "Kullanıcı", "Alıcı", "Konu", "Durum", "Hata", "Gönderen"],
        rows=excel_rows,
        widths={"A": 22, "B": 22, "C": 28, "D": 30, "E": 48, "F": 14, "G": 36, "H": 28},
    )
    return build_excel_download_response(output, download_name=f"bys360_performans_mail_gecmisi_{period_id}.xlsx")

def build_performance_report_excel_download_response(evaluations: Iterable[Any], *, download_name: str = "bys360_performans_raporu.xlsx"):
    rows: list[list[Any]] = []
    for item in evaluations:
        employee = getattr(item, "employee", None)
        period = getattr(item, "period", None)
        employee_name = _full_name(employee)
        rows.append([
            getattr(period, "title", None) or "-",
            employee_name,
            getattr(employee, "sicil_no", None) or "-",
            getattr(employee, "unvan", None) or "-",
            getattr(employee, "birim", None) or "-",
            getattr(employee, "ust_birim", None) or "-",
            float(getattr(item, "level_1_total_100", 0) or 0),
            float(getattr(item, "level_2_total_100", 0) or 0),
            float(getattr(item, "level_3_total_100", 0) or 0),
            float(getattr(item, "report_final_score", 0) or 0),
            getattr(item, "status", None) or "-",
            "Evet" if bool(getattr(period, "results_published", False)) else "Hayır",
        ])

    output = build_styled_excel_bytes(
        title="Performans Raporu",
        headers=[
            "Dönem",
            "Personel",
            "Sicil No",
            "Unvan",
            "Birim",
            "Üst Birim",
            "1. Amir Puanı",
            "2. Amir Puanı",
            "3. Amir Puanı",
            "Nihai Puan",
            "Durum",
            "Sonuç Yayını",
        ],
        rows=rows,
        widths={
            "A": 28,
            "B": 28,
            "C": 16,
            "D": 22,
            "E": 24,
            "F": 24,
            "G": 14,
            "H": 14,
            "I": 14,
            "J": 14,
            "K": 18,
            "L": 16,
        },
    )
    return build_excel_download_response(output, download_name=download_name)

__all__ = [
    "CSV_MIMETYPE",
    "XLSX_MIMETYPE",
    "build_binary_download_response",
    "build_csv_text_download_response",
    "build_excel_download_response",
    "build_mail_history_export_response",
    "build_performance_report_excel_download_response",
    "build_publish_log_export_response",
    "build_styled_excel_bytes",
]