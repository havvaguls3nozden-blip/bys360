
"""Analiz Merkezi Faz 7: güvenli Excel/CSV ön izleme servisi.

Bu servis gerçek içe aktarım yapmaz. Dosya diske yazılmaz, veritabanına kayıt
atılmaz ve yalnızca istek süresince bellekte okunur. Amaç; Excel/CSV dosyası
kurumsal raporlama hattına alınmadan önce güvenli dosya doğrulama, sütun
analizi, örnek satır önizleme ve KVKK uyarısı üretmektir.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any
from collections.abc import Iterable, Sequence

from werkzeug.datastructures import FileStorage

from app.security.upload_security import UploadValidationError, validate_upload

GERCEK_ICE_AKTARIM_YOK = True
DB_WRITE_ENABLED = False
ALLOWED_ANALYSIS_EXTENSIONS = {"xlsx", "csv"}
DEFAULT_SAMPLE_LIMIT = 8
DEFAULT_SCAN_LIMIT = 250
MAX_ANALYSIS_FILE_SIZE = 8 * 1024 * 1024
MAX_XLSX_ZIP_ENTRIES = 420
MAX_XLSX_UNCOMPRESSED_BYTES = 28 * 1024 * 1024
MAX_CELL_TEXT = 160

_HIGH_RISK_HEADER_PATTERNS: tuple[str, ...] = (
    "tc kimlik",
    "tckn",
    "kimlik no",
    "kimlik",
    "iban",
    "banka",
    "maas",
    "maaş",
    "ucret",
    "ücret",
    "saglik",
    "sağlık",
    "hastalik",
    "hastalık",
    "engel",
    "rapor",
    "adres",
    "telefon",
    "gsm",
    "cep",
    "mail",
    "e posta",
    "eposta",
    "email",
    "dogum",
    "doğum",
    "disiplin",
    "ceza",
)
_MEDIUM_RISK_HEADER_PATTERNS: tuple[str, ...] = (
    "sicil",
    "ad",
    "soyad",
    "isim",
    "personel",
    "unvan",
    "birim",
    "amir",
    "yonetici",
    "yönetici",
    "performans",
    "puan",
    "not",
    "izin",
    "devamsizlik",
    "devamsızlık",
    "anket",
    "geri bildirim",
)
_FORMULA_PREFIXES = ("=", "+", "-", "@")
_EMAIL_RE = re.compile(r"^[^@\s]{2,}@[^@\s]+\.[^@\s]+$", re.IGNORECASE)
_TCKN_RE = re.compile(r"(?<!\d)\d{11}(?!\d)")
_PHONE_RE = re.compile(r"(?<!\d)(?:\+?90)?0?5\d{9}(?!\d)")
_IBAN_RE = re.compile(r"\bTR\d{2}[A-Z0-9]{5,}\b", re.IGNORECASE)
_LONG_DIGIT_RE = re.compile(r"\d{7,}")


@dataclass(slots=True)
class ExcelPreviewValidationError(ValueError):
    message: str

    def __str__(self) -> str:  # pragma: no cover - dataclass convenience
        return self.message


@dataclass(slots=True)
class ColumnProfile:
    index: int
    header: str
    normalized_header: str
    inferred_type: str = "boş"
    sensitivity: str = "düşük"
    kvkk_reason: str = "Kişisel veri sinyali düşük."
    non_empty_count: int = 0
    empty_count: int = 0
    formula_count: int = 0
    sample_values: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "index": self.index,
            "header": self.header,
            "normalized_header": self.normalized_header,
            "inferred_type": self.inferred_type,
            "sensitivity": self.sensitivity,
            "kvkk_reason": self.kvkk_reason,
            "non_empty_count": self.non_empty_count,
            "empty_count": self.empty_count,
            "formula_count": self.formula_count,
            "sample_values": self.sample_values,
        }


@dataclass(slots=True)
class SafeExcelPreviewResult:
    ok: bool
    import_enabled: bool
    db_write_enabled: bool
    original_filename: str
    extension: str
    detected_mime: str
    file_size: int
    sheet_name: str
    sheet_names: list[str]
    header_row_index: int
    row_count_visible: int
    row_count_estimated: int
    column_count: int
    preview_limit: int
    scan_limit: int
    truncated: bool
    columns: list[ColumnProfile]
    preview_rows: list[dict[str, Any]]
    warnings: list[str]
    kvkk_warnings: list[str]
    security_findings: list[str]
    duplicate_headers: list[str]
    empty_headers: list[int]
    formula_cells: int

    @property
    def sensitive_column_count(self) -> int:
        return sum(1 for column in self.columns if column.sensitivity in {"orta", "yüksek"})

    @property
    def metrics(self) -> dict[str, Any]:
        return {
            "row_count_visible": self.row_count_visible,
            "row_count_estimated": self.row_count_estimated,
            "column_count": self.column_count,
            "sensitive_column_count": self.sensitive_column_count,
            "formula_cells": self.formula_cells,
            "duplicate_header_count": len(self.duplicate_headers),
            "empty_header_count": len(self.empty_headers),
            "import_enabled": self.import_enabled,
            "db_write_enabled": self.db_write_enabled,
        }

    def to_dict(self) -> dict[str, Any]:
        return {
            "ok": self.ok,
            "import_enabled": self.import_enabled,
            "db_write_enabled": self.db_write_enabled,
            "original_filename": self.original_filename,
            "extension": self.extension,
            "detected_mime": self.detected_mime,
            "file_size": self.file_size,
            "sheet_name": self.sheet_name,
            "sheet_names": self.sheet_names,
            "header_row_index": self.header_row_index,
            "row_count_visible": self.row_count_visible,
            "row_count_estimated": self.row_count_estimated,
            "column_count": self.column_count,
            "preview_limit": self.preview_limit,
            "scan_limit": self.scan_limit,
            "truncated": self.truncated,
            "columns": [column.to_dict() for column in self.columns],
            "preview_rows": self.preview_rows,
            "warnings": self.warnings,
            "kvkk_warnings": self.kvkk_warnings,
            "security_findings": self.security_findings,
            "duplicate_headers": self.duplicate_headers,
            "empty_headers": self.empty_headers,
            "formula_cells": self.formula_cells,
            "metrics": self.metrics,
        }


def _safe_int(value: Any, default: int, minimum: int, maximum: int) -> int:
    try:
        parsed = int(str(value or default).strip())
    except (TypeError, ValueError):
        parsed = default
    return max(min(parsed, maximum), minimum)


def normalize_header(value: Any) -> str:
    text = _cell_to_text(value).strip().lower()
    translation = str.maketrans("çğıöşüÇĞİÖŞÜ", "cgiosuCGIOSU")
    text = text.translate(translation)
    text = re.sub(r"[_\-./]+", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if len(text) > MAX_CELL_TEXT:
        return text[:MAX_CELL_TEXT] + "…"
    return text


def _is_formula_text(value: Any) -> bool:
    text = str(value or "").lstrip()
    if not text:
        return False
    if text.startswith("="):
        return True
    if text[0] in _FORMULA_PREFIXES and len(text) > 1 and re.search(r"[A-Za-z][0-9]|SUM|ORTALAMA|DÜŞEYARA|VLOOKUP", text, re.IGNORECASE):
        return True
    return False


def _header_sensitivity(normalized_header: str) -> tuple[str, str]:
    if any(pattern in normalized_header for pattern in _HIGH_RISK_HEADER_PATTERNS):
        return "yüksek", "Başlık kişisel veya özel nitelikli veri alanı gibi görünüyor."
    if any(pattern in normalized_header for pattern in _MEDIUM_RISK_HEADER_PATTERNS):
        return "orta", "Başlık kurumsal/personel verisi içerebilir; görünürlük sınırlı tutulmalı."
    return "düşük", "Kişisel veri sinyali düşük."


def _value_sensitivity(text: str) -> tuple[str | None, str | None]:
    if not text:
        return None, None
    if _IBAN_RE.search(text):
        return "yüksek", "IBAN benzeri finansal veri algılandı."
    if _TCKN_RE.search(text):
        return "yüksek", "TC kimlik numarası benzeri 11 haneli veri algılandı."
    if _EMAIL_RE.match(text):
        return "yüksek", "E-posta adresi algılandı."
    if _PHONE_RE.search(text):
        return "yüksek", "Telefon numarası benzeri veri algılandı."
    if _LONG_DIGIT_RE.search(text):
        return "orta", "Uzun sayısal tanımlayıcı algılandı."
    return None, None


def _merge_sensitivity(header_level: str, value_level: str | None) -> str:
    levels = {"düşük": 0, "orta": 1, "yüksek": 2, None: -1}
    merged = max(levels.get(header_level, 0), levels.get(value_level, -1))
    return "yüksek" if merged >= 2 else "orta" if merged == 1 else "düşük"


def mask_cell_value(value: Any, header: str = "") -> str:
    """KVKK amaçlı ekran maskelemesi; veri yazmaz, yalnız görünümü dönüştürür."""
    if _is_formula_text(value):
        return "[FORMÜL GİZLENDİ]"
    text = _cell_to_text(value)
    if not text:
        return ""
    normalized = normalize_header(header)

    if _IBAN_RE.search(text):
        return _IBAN_RE.sub(lambda match: match.group(0)[:6] + "…" + match.group(0)[-4:], text)
    if _EMAIL_RE.match(text):
        local, domain = text.split("@", 1)
        return f"{local[:1]}***@{domain}"
    if _TCKN_RE.fullmatch(text):
        return f"{text[:2]}*******{text[-2:]}"
    if _PHONE_RE.fullmatch(text):
        return f"{text[:3]}******{text[-2:]}"
    if "sicil" in normalized and len(text) > 3:
        return f"***{text[-3:]}"
    if any(key in normalized for key in ("ad", "soyad", "isim", "personel")):
        pieces = [piece for piece in re.split(r"\s+", text) if piece]
        if pieces and len(text) > 2:
            return " ".join(piece[:1] + "***" for piece in pieces)
    if any(key in normalized for key in ("adres", "saglik", "hastalik", "engel", "maas", "ucret", "banka")):
        return "[KVKK nedeniyle maskelendi]"
    return _LONG_DIGIT_RE.sub(lambda match: match.group(0)[:2] + "…" + match.group(0)[-2:], text)


def _infer_type(values: Iterable[Any]) -> str:
    counts: Counter[str] = Counter()
    for raw in values:
        text = _cell_to_text(raw)
        if not text:
            continue
        if _is_formula_text(raw):
            counts["formül"] += 1
        elif _EMAIL_RE.match(text):
            counts["e-posta"] += 1
        elif _IBAN_RE.search(text):
            counts["iban/finans"] += 1
        elif _TCKN_RE.fullmatch(text):
            counts["kimlik no"] += 1
        elif _PHONE_RE.fullmatch(text):
            counts["telefon"] += 1
        else:
            try:
                float(text.replace(",", "."))
                counts["sayısal"] += 1
            except ValueError:
                counts["metin"] += 1
    if not counts:
        return "boş"
    return counts.most_common(1)[0][0]


def _first_non_empty_row(rows: Sequence[Sequence[Any]]) -> tuple[int, list[Any]]:
    for index, row in enumerate(rows, start=1):
        if any(_cell_to_text(cell) for cell in row):
            return index, list(row)
    raise ExcelPreviewValidationError("Dosyada okunabilir başlık satırı bulunamadı.")


def _normalize_matrix_width(headers: list[Any], rows: list[list[Any]]) -> tuple[list[Any], list[list[Any]]]:
    width = max([len(headers), *(len(row) for row in rows)] or [0])
    if width <= 0:
        raise ExcelPreviewValidationError("Dosyada okunabilir sütun bulunamadı.")
    normalized_headers = list(headers) + [""] * (width - len(headers))
    normalized_rows = [list(row) + [None] * (width - len(row)) for row in rows]
    return normalized_headers[:width], [row[:width] for row in normalized_rows]


def _build_profiles(headers: list[Any], rows: list[list[Any]]) -> tuple[list[ColumnProfile], list[str], list[int], int]:
    normalized_headers = [normalize_header(header) for header in headers]
    header_counts = Counter(h for h in normalized_headers if h)
    duplicate_headers = [header for header, count in header_counts.items() if count > 1]
    empty_headers = [idx + 1 for idx, header in enumerate(normalized_headers) if not header]
    profiles: list[ColumnProfile] = []
    total_formula_cells = 0

    for col_idx, raw_header in enumerate(headers):
        header = _cell_to_text(raw_header) or f"Kolon {col_idx + 1}"
        normalized_header = normalized_headers[col_idx]
        header_level, reason = _header_sensitivity(normalized_header)
        values = [row[col_idx] for row in rows]
        non_empty_values = [value for value in values if _cell_to_text(value)]
        formula_count = sum(1 for value in values if _is_formula_text(value))
        total_formula_cells += formula_count
        sensitivity = header_level
        kvkk_reason = reason
        sample_values: list[str] = []
        for value in non_empty_values[:4]:
            text = _cell_to_text(value)
            value_level, value_reason = _value_sensitivity(text)
            sensitivity = _merge_sensitivity(sensitivity, value_level)
            if value_reason and sensitivity == "yüksek":
                kvkk_reason = value_reason
            masked = mask_cell_value(value, header)
            if masked not in sample_values:
                sample_values.append(masked)
        profiles.append(
            ColumnProfile(
                index=col_idx + 1,
                header=header,
                normalized_header=normalized_header or f"kolon {col_idx + 1}",
                inferred_type=_infer_type(values),
                sensitivity=sensitivity,
                kvkk_reason=kvkk_reason,
                non_empty_count=len(non_empty_values),
                empty_count=max(len(rows) - len(non_empty_values), 0),
                formula_count=formula_count,
                sample_values=sample_values,
            )
        )
    return profiles, duplicate_headers, empty_headers, total_formula_cells


def _build_preview_rows(headers: list[Any], rows: list[list[Any]], sample_limit: int) -> list[dict[str, Any]]:
    preview_rows: list[dict[str, Any]] = []
    for row_number, row in enumerate(rows[:sample_limit], start=1):
        cells: list[dict[str, Any]] = []
        for idx, value in enumerate(row):
            header = _cell_to_text(headers[idx]) or f"Kolon {idx + 1}"
            text = _cell_to_text(value)
            value_level, value_reason = _value_sensitivity(text)
            header_level, _ = _header_sensitivity(normalize_header(header))
            cells.append(
                {
                    "index": idx + 1,
                    "header": header,
                    "value": mask_cell_value(value, header),
                    "empty": not bool(text),
                    "formula": _is_formula_text(value),
                    "sensitivity": _merge_sensitivity(header_level, value_level),
                    "reason": value_reason or "",
                }
            )
        preview_rows.append({"row_number": row_number, "cells": cells})
    return preview_rows


def _security_scan_xlsx(payload: bytes) -> list[str]:
    findings: list[str] = []
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_XLSX_ZIP_ENTRIES:
                raise ExcelPreviewValidationError("Excel dosyası beklenenden fazla iç parça içeriyor; güvenlik nedeniyle okunmadı.")
            total_uncompressed = sum(info.file_size for info in infos)
            if total_uncompressed > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ExcelPreviewValidationError("Excel dosyasının açılmış boyutu güvenli sınırı aşıyor; önizleme yapılmadı.")
            names = [info.filename.lower() for info in infos]
            if any("vbaproject.bin" in name for name in names):
                findings.append("Makro/VBA izi algılandı. .xlsx dosyasında beklenmeyen makro parçası var; gerçek aktarım yapılmamalı.")
            if any(name.startswith("xl/externallinks/") for name in names):
                findings.append("Dış bağlantı izi algılandı. Dosya harici kaynak referansı içerebilir.")
            if any("activex" in name or "embeddings" in name or "oleobjects" in name for name in names):
                findings.append("Gömülü nesne/ActiveX izi algılandı. Yönetici kontrolü olmadan kullanılmamalı.")
            if any(name.endswith(".rels") for name in names):
                for name in names:
                    if not name.endswith(".rels"):
                        continue
                    try:
                        rel_text = archive.read(name).decode("utf-8", errors="ignore")
                    except Exception:
                        __import__("logging").getLogger(__name__).exception("BYS360 kalite denetimi: except bloğu loglandı (app/services/ai/excel_preview.py:444)")
                        continue
                    if "TargetMode=\"External\"" in rel_text or "TargetMode='External'" in rel_text:
                        findings.append("Dosyada harici ilişki/referans izi algılandı.")
                        break
    except zipfile.BadZipFile as exc:
        raise ExcelPreviewValidationError("Excel dosyası geçerli bir .xlsx paketi gibi okunamadı.") from exc
    return findings


def _decode_csv(payload: bytes) -> tuple[str, str]:
    for encoding in ("utf-8-sig", "cp1254", "iso-8859-9", "latin-1"):
        try:
            return payload.decode(encoding), encoding
        except UnicodeDecodeError:
            __import__("logging").getLogger(__name__).exception("BYS360 kalite denetimi: except bloğu loglandı (app/services/ai/excel_preview.py:458)")
            continue
    raise ExcelPreviewValidationError("CSV dosya kodlaması okunamadı.")


def _read_csv_rows(payload: bytes, scan_limit: int) -> tuple[list[Any], list[list[Any]], list[str], int, int, bool]:
    text, encoding = _decode_csv(payload)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,|\t,")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text, newline=""), dialect)
    raw_rows: list[list[Any]] = []
    estimated = 0
    truncated = False
    for row in reader:
        estimated += 1
        if len(raw_rows) <= scan_limit + 1:
            raw_rows.append(row)
        else:
            truncated = True
            break
    if not raw_rows:
        raise ExcelPreviewValidationError("CSV dosyasında okunabilir satır bulunamadı.")
    header_index, headers = _first_non_empty_row(raw_rows)
    data_rows = raw_rows[header_index: header_index + scan_limit]
    warnings = [f"CSV kodlaması {encoding} olarak okundu."]
    return headers, data_rows, warnings, header_index, estimated, truncated


def _read_xlsx_rows(payload: bytes, scan_limit: int, sheet_name: str | None = None) -> tuple[list[Any], list[list[Any]], list[str], str, list[str], int, int, bool]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover - deployment dependency guard
        raise ExcelPreviewValidationError("openpyxl paketi bulunamadı; .xlsx önizleme için requirements kurulmalı.") from exc

    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=False, keep_links=False)
    sheet_names = list(workbook.sheetnames)
    if not sheet_names:
        raise ExcelPreviewValidationError("Excel dosyasında çalışma sayfası bulunamadı.")
    selected_sheet = (sheet_name or "").strip()
    if selected_sheet and selected_sheet in sheet_names:
        worksheet = workbook[selected_sheet]
    else:
        worksheet = workbook[sheet_names[0]]
    selected_sheet = worksheet.title

    raw_rows: list[list[Any]] = []
    estimated = int(worksheet.max_row or 0)
    truncated = False
    for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        raw_rows.append(list(row))
        if row_index > scan_limit + 1:
            truncated = True
            break
    if not raw_rows:
        raise ExcelPreviewValidationError("Excel dosyasında okunabilir satır bulunamadı.")
    header_index, headers = _first_non_empty_row(raw_rows)
    data_rows = raw_rows[header_index: header_index + scan_limit]
    warnings = []
    if len(sheet_names) > 1:
        warnings.append("Dosyada birden fazla sayfa var. Önizleme seçili sayfa üzerinden yapılır; sayfa değişimi için dosyayı yeniden yükleyin.")
    if worksheet.max_column and worksheet.max_column > 120:
        warnings.append("Sütun sayısı yüksek görünüyor; gerçek raporlama öncesi dosya sadeleştirilmeli.")
    return headers, data_rows, warnings, selected_sheet, sheet_names, header_index, estimated, truncated


def build_safe_excel_preview(
    file: FileStorage,
    *,
    sample_limit: int | str | None = DEFAULT_SAMPLE_LIMIT,
    scan_limit: int | str | None = DEFAULT_SCAN_LIMIT,
    sheet_name: str | None = None,
) -> SafeExcelPreviewResult:
    """Validate and preview an Excel/CSV file without importing or persisting data."""
    sample_limit_int = _safe_int(sample_limit, DEFAULT_SAMPLE_LIMIT, 3, 25)
    scan_limit_int = _safe_int(scan_limit, DEFAULT_SCAN_LIMIT, 50, 1000)
    try:
        meta = validate_upload(
            file,
            allowed_extensions=ALLOWED_ANALYSIS_EXTENSIONS,
            max_size=MAX_ANALYSIS_FILE_SIZE,
        )
    except UploadValidationError as exc:
        raise ExcelPreviewValidationError(str(exc)) from exc

    extension = str(meta.get("extension") or Path(file.filename or "").suffix.lower().lstrip("."))
    if extension not in ALLOWED_ANALYSIS_EXTENSIONS:
        raise ExcelPreviewValidationError("Bu fazda yalnız .xlsx ve .csv dosyaları güvenli önizlemeye alınır.")

    file.stream.seek(0)
    payload = file.stream.read(MAX_ANALYSIS_FILE_SIZE + 1)
    file.stream.seek(0)
    if len(payload) > MAX_ANALYSIS_FILE_SIZE:
        raise ExcelPreviewValidationError("Dosya boyutu güvenli analiz sınırını aşıyor.")

    security_findings: list[str] = []
    warnings: list[str] = [
        "Gerçek içe aktarım yapılmaz; bu ekran yalnız güvenli önizleme üretir.",
        "Dosya kalıcı olarak kaydedilmez ve veritabanına yazılmaz.",
    ]
    selected_sheet = "CSV"
    sheet_names: list[str] = []
    if extension == "xlsx":
        security_findings.extend(_security_scan_xlsx(payload))
        headers, data_rows, read_warnings, selected_sheet, sheet_names, header_row_index, estimated_rows, truncated = _read_xlsx_rows(
            payload,
            scan_limit_int,
            sheet_name=sheet_name,
        )
    else:
        headers, data_rows, read_warnings, header_row_index, estimated_rows, truncated = _read_csv_rows(payload, scan_limit_int)
        sheet_names = ["CSV"]
    warnings.extend(read_warnings)

    headers, data_rows = _normalize_matrix_width(list(headers), data_rows)
    columns, duplicate_headers, empty_headers, formula_cells = _build_profiles(headers, data_rows)
    preview_rows = _build_preview_rows(headers, data_rows, sample_limit_int)

    if duplicate_headers:
        warnings.append("Tekrarlanan sütun başlığı var: " + ", ".join(duplicate_headers[:8]))
    if empty_headers:
        warnings.append("Boş başlıklı sütunlar var: " + ", ".join(str(idx) for idx in empty_headers[:12]))
    if formula_cells:
        warnings.append(f"{formula_cells} formül hücresi algılandı; formüller çalıştırılmadan gizli gösterildi.")
    if truncated or estimated_rows > len(data_rows) + 1:
        warnings.append("Önizleme güvenli sınır nedeniyle örnek satırlarla sınırlandırıldı.")

    kvkk_warnings = _build_kvkk_warnings(columns)
    return SafeExcelPreviewResult(
        ok=True,
        import_enabled=False,
        db_write_enabled=False,
        original_filename=str(meta.get("original_filename") or file.filename or ""),
        extension=extension,
        detected_mime=str(meta.get("detected_mime") or ""),
        file_size=int(meta.get("size") or len(payload)),
        sheet_name=selected_sheet,
        sheet_names=sheet_names,
        header_row_index=header_row_index,
        row_count_visible=len(data_rows),
        row_count_estimated=max(estimated_rows - header_row_index, len(data_rows)),
        column_count=len(headers),
        preview_limit=sample_limit_int,
        scan_limit=scan_limit_int,
        truncated=truncated,
        columns=columns,
        preview_rows=preview_rows,
        warnings=warnings,
        kvkk_warnings=kvkk_warnings,
        security_findings=security_findings,
        duplicate_headers=duplicate_headers,
        empty_headers=empty_headers,
        formula_cells=formula_cells,
    )


def _build_kvkk_warnings(columns: Sequence[ColumnProfile]) -> list[str]:
    high = [column.header for column in columns if column.sensitivity == "yüksek"]
    medium = [column.header for column in columns if column.sensitivity == "orta"]
    warnings: list[str] = []
    if high:
        warnings.append("Yüksek hassasiyetli kişisel veri sütunları algılandı: " + ", ".join(high[:10]))
    if medium:
        warnings.append("Personel/kurumsal veri içerebilecek sütunlar algılandı: " + ", ".join(medium[:10]))
    if high or medium:
        warnings.append("KVKK uyarısı: Önizleme ekranı değerleri maskeler; gerçek aktarım ayrı onay, amaç sınırlılığı ve yetki kontrolü gerektirir.")
    else:
        warnings.append("KVKK uyarısı: Belirgin kişisel veri sütunu algılanmadı; yine de dosya içeriği kurumsal gizlilik kapsamında değerlendirilmelidir.")
    return warnings


def build_empty_excel_preview_context() -> dict[str, Any]:
    return {
        "preview": None,
        "phase_label": "Faz 7",
        "import_enabled": False,
        "db_write_enabled": False,
        "allowed_extensions": sorted(ALLOWED_ANALYSIS_EXTENSIONS),
        "max_file_size_mb": MAX_ANALYSIS_FILE_SIZE // (1024 * 1024),
        "default_sample_limit": DEFAULT_SAMPLE_LIMIT,
        "default_scan_limit": DEFAULT_SCAN_LIMIT,
        "safety_notes": [
            "Gerçek içe aktarım yapılmaz.",
            "Dosya sunucuya kalıcı olarak kaydedilmez.",
            "Örnek satırlar KVKK maskesiyle gösterilir.",
            "Sütun analizi karar destek amaçlıdır; idari karar üretmez.",
        ],
    }
