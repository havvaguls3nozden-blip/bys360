"""Operational admin user import service extracted from ops_routes.py.

Route decorators stay in ops_routes.py. This module contains the heavy implementation body.
"""
from __future__ import annotations

from app.services.safe_user_delete_service import safe_delete_user_by_id
from app.core.datetime_utils import utc_now
from flask import current_app, flash, redirect, request, send_file, url_for
from flask_login import current_user, login_required
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError
from app.extensions import db
from app.models import EmployeeOrgAssignmentHistory, OrganizationUnit, PerformancePeriod, User
from app.route_registry import main_bp
from app.route_support import admin_required, menu_key_required, safe_db_rollback, safe_render, ensure_boolean_toggle, normalize_int_list
from app.services.hierarchy_admin_service import normalize_text, reset_all_personnel_and_related_data
from app.services.auto_hierarchy_service import auto_apply_manager_chains, infer_role_from_profile
from app.services.personnel_sync_service import canonical_role_label, canonical_role_value
from app.services.personnel.categories import assign_user_performance_category, normalize_personnel_category_label
from app.services.performance_service import generate_assignments_for_active_period, is_single_manager_case
from app.services.ai import build_personnel_profile_chain_ai_panel, build_excel_fix_preview_ai_panel, build_import_health_priority_ai_panel, build_import_health_simulation_ai_panel
from app.services.personnel.import_manager_chain_sync import sync_touched_users_manager_ids_from_sicils
from app.services.personnel.excel_import_guard import validate_personnel_import_rows_for_commit
from app.services.profile_photo_service import (
    delete_profile_photo_file as _delete_profile_photo_file,
    save_profile_photo as _save_profile_photo,
)
from .ops_helpers import (
    _profile_full_name,
    _safe_text,
    _resolve_user_by_sicil,
    _build_org_path,
    _serialize_manager_card,
    _resolve_personnel_profile_hierarchy,
    _canonicalize_import_headers,
    _collapse_spaces,
    _has_explicit_manager_columns,
    get_default_first_login_password,
)
from .ops_health_services import admin_import_health_report_impl

def admin_user_import_impl():
    if request.method == "POST":
        f = request.files.get("excel_file")
        if not f or not f.filename:
            flash("Lütfen bir Excel dosyası seçin.", "warning")
            return safe_render("excel_import.html", "<h3>Excel İçe Aktar</h3>", apply_ai_fixes=False, ai_excel_fix_panel=build_excel_fix_preview_ai_panel())

        wb = load_workbook(f, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            flash("Excel dosyası boş.", "danger")
            return safe_render("excel_import.html", "<h3>Excel İçe Aktar</h3>", apply_ai_fixes=False, ai_excel_fix_panel=build_excel_fix_preview_ai_panel())

        import_preflight = validate_personnel_import_rows_for_commit(rows)
        if not import_preflight.get("ok"):
            errors = list(import_preflight.get("errors") or [])
            warnings = list(import_preflight.get("warnings") or [])
            flash("Excel ön kontrolü başarısız. Hiçbir kayıt veritabanına yazılmadı.", "danger")
            return safe_render(
                "excel_import.html",
                "<h3>Excel İçe Aktar</h3>",
                errors=errors,
                import_preflight=import_preflight,
                import_warnings=warnings,
                apply_ai_fixes=False,
                ai_excel_fix_panel=build_excel_fix_preview_ai_panel(errors=errors, normalized_headers=import_preflight.get("canonical_headers") or []),
            )

        apply_ai_fixes = (request.form.get("apply_ai_fixes") or "").strip().lower() in {"1", "true", "on", "evet", "yes"}
        headers = [normalize_text(h) if h is not None else "" for h in rows[0]]
        normalized_headers, header_alias_count = _canonicalize_import_headers(headers, apply_ai_fixes)
        autofix_stats = {"header_alias_count": header_alias_count, "whitespace_trim_count": 0, "inferred_role_count": 0, "default_active_count": 0}
        idx = {h: i for i, h in enumerate(normalized_headers)}

        def getv(row, name, default=""):
            key = normalize_text(name)
            if key not in idx:
                return default
            value = row[idx[key]]
            if value is None:
                return default
            if isinstance(value, float) and value.is_integer():
                return str(int(value))
            if isinstance(value, int):
                return str(value)
            return str(value).strip()

        def getv_any(row, *names, default=""):
            for name in names:
                value = getv(row, name, None)
                if value is not None and str(value).strip() != "":
                    return str(value).strip()
            return default

        created = 0
        updated = 0
        skipped_admin = 0
        errors = []
        touched_users = []

        for rowno, row in enumerate(rows[1:], start=2):
            if not row or all(v is None or str(v).strip() == "" for v in row):
                continue

            sicil_no = _collapse_spaces(getv_any(row, "sicil_no", "sicil no", "sicil"), apply_ai_fixes)
            ad = _collapse_spaces(getv_any(row, "ad", "adi", "adı"), apply_ai_fixes)
            soyad = _collapse_spaces(getv_any(row, "soyad", "soyadi", "soyadı"), apply_ai_fixes)
            email = _collapse_spaces(getv_any(row, "email", "e-posta", "eposta", default="").lower(), apply_ai_fixes)
            unvan = _collapse_spaces(getv_any(row, "unvan", "ünvan"), apply_ai_fixes)
            birim = _collapse_spaces(getv_any(row, "birim"), apply_ai_fixes)
            ust_birim = _collapse_spaces(getv_any(row, "ust_birim", "ust birim", "üst birim"), apply_ai_fixes)
            raw_role = getv_any(row, "role", "rol", default="")
            inferred_role, inferred_role_label = infer_role_from_profile(raw_role=raw_role, unvan=unvan, birim=birim, ust_birim=ust_birim)
            if apply_ai_fixes and not normalize_text(raw_role) and inferred_role:
                autofix_stats["inferred_role_count"] += 1
            role = canonical_role_value(raw_role or inferred_role or "personel")
            role_label = canonical_role_label(raw_role or inferred_role_label or role)
            yonetici_sicil = _collapse_spaces(getv_any(
                row,
                "yonetici_sicil",
                "yonetici sicil",
                "yönetici sicil",
                "1. amir sicil",
                "1 amir sicil",
                "birinci amir sicil",
                "1. yönetici sicil",
                "1 yonetici sicil",
                "1. amir",
                "1 amir",
                "birinci amir",
            ), apply_ai_fixes)
            ikinci_yonetici_sicil = _collapse_spaces(getv_any(
                row,
                "ikinci_yonetici_sicil",
                "ikinci yonetici sicil",
                "ikinci yönetici sicil",
                "2. amir sicil",
                "2 amir sicil",
                "ikinci amir sicil",
                "2. yönetici sicil",
                "2 yonetici sicil",
                "2. amir",
                "2 amir",
                "ikinci amir",
            ), apply_ai_fixes)
            ucuncu_yonetici_sicil = _collapse_spaces(getv_any(
                row,
                "ucuncu_yonetici_sicil",
                "ucuncu yonetici sicil",
                "üçüncü yönetici sicil",
                "ucuncu yönetici sicil",
                "3. amir sicil",
                "3 amir sicil",
                "üçüncü amir sicil",
                "ucuncu amir sicil",
                "3. yönetici sicil",
                "3 yonetici sicil",
                "3. amir",
                "3 amir",
                "üçüncü amir",
                "ucuncu amir",
            ), apply_ai_fixes)
            personnel_category = normalize_personnel_category_label(getv_any(
                row,
                "personnel_category",
                "kategori",
                "personel kategori",
                "personel kategorisi",
                "performans kategori",
                "performans kategorisi",
                default="Diğer",
            ))  # BYS360_PHASE2_3_IMPORT_CATEGORY_READ
            is_active_source = getv_any(row, "is_active", "aktif", default="1")
            if apply_ai_fixes and not str(is_active_source or "").strip():
                autofix_stats["default_active_count"] += 1
            is_active_raw = normalize_text(is_active_source)
            is_active = is_active_raw in {"1", "true", "evet", "aktif", "yes", ""}

            if apply_ai_fixes and any("  " in str(v or "") for v in row):
                autofix_stats["whitespace_trim_count"] += 1
            if not sicil_no or not ad or not soyad or not email or not unvan or not birim:
                errors.append(f"Satır {rowno}: zorunlu alan eksik.")
                continue
            if normalize_text(role) == "admin":
                skipped_admin += 1
                continue

            unit = None
            try:
                # Birim otomatik oluşsun; bu ekrandaki en kritik işlerden biriydi.
                from app.services.hierarchy_admin_service import ensure_unit_exists_strict
                unit = ensure_unit_exists_strict(birim, ust_birim, role)
            except Exception as exc:
                __import__("logging").getLogger(__name__).exception("BYS360 kalite denetimi: except bloğu loglandı (app/admin/ops_routes.py:621)")
                errors.append(f"Satır {rowno}: birim oluşturulamadı -> {exc}")
                continue

            if not unit:
                errors.append(f"Satır {rowno}: birim çözümlenemedi -> {birim} / {ust_birim}")
                continue

            user = User.query.filter_by(sicil_no=sicil_no).first()
            if user and user.role == "admin":
                skipped_admin += 1
                continue

            resolved_ust_birim = ust_birim if ust_birim else (unit.parent.name if unit.parent else unit.name)

            try:
                if user:
                    user.ad = ad
                    user.soyad = soyad
                    user.email = email
                    user.unvan = unvan
                    user.role = role or "personel"
                    if hasattr(user, "role_label"):
                        user.role_label = role_label
                    user.organization_unit_id = unit.id
                    user.birim = birim
                    user.ust_birim = resolved_ust_birim
                    user.yonetici_sicil = yonetici_sicil or None
                    user.ikinci_yonetici_sicil = ikinci_yonetici_sicil or None
                    user.ucuncu_yonetici_sicil = ucuncu_yonetici_sicil or None
                    user.is_active = is_active
                    assign_user_performance_category(user, personnel_category, db_session=db.session)  # BYS360_PHASE2_3_IMPORT_CATEGORY_UPDATE_BINDING
                    updated += 1
                    touched_users.append(user)
                else:
                    user = User(
                        ad=ad,
                        soyad=soyad,
                        sicil_no=sicil_no,
                        email=email,
                        unvan=unvan,
                        role=role or "personel",
                        role_label=role_label,
                        organization_unit_id=unit.id,
                        birim=birim,
                        ust_birim=resolved_ust_birim,
                        yonetici_sicil=yonetici_sicil or None,
                        ikinci_yonetici_sicil=ikinci_yonetici_sicil or None,
                        ucuncu_yonetici_sicil=ucuncu_yonetici_sicil or None,
                        is_active=is_active,
                    )
                    assign_user_performance_category(user, personnel_category, db_session=db.session)  # BYS360_PHASE2_3_IMPORT_CATEGORY_CREATE_BINDING
                    user.set_password(get_default_first_login_password())
                    if hasattr(user, "must_change_password"):
                        user.must_change_password = True
                    if hasattr(user, "must_set_security_question"):
                        user.must_set_security_question = True
                    if hasattr(user, "is_first_login"):
                        user.is_first_login = True
                    db.session.add(user)
                    created += 1
                    touched_users.append(user)
            except Exception as exc:
                errors.append(f"Satır {rowno}: {exc}")

        if errors:
            db.session.rollback()
            flash("Excel içe aktarma sırasında hatalar oluştu.", "danger")
            return safe_render("excel_import.html", "<h3>Excel İçe Aktar</h3>", errors=errors, apply_ai_fixes=apply_ai_fixes, ai_excel_fix_panel=build_excel_fix_preview_ai_panel(errors=errors, normalized_headers=normalized_headers, autofix_enabled=apply_ai_fixes, applied_fixes=autofix_stats))

        db.session.flush()
        has_explicit_manager_columns = _has_explicit_manager_columns(normalized_headers)
        if has_explicit_manager_columns:
            # BYS360_IMPORT_EXPLICIT_MANAGER_CHAIN_PRESERVE_V1
            manager_sync_summary = sync_touched_users_manager_ids_from_sicils(
                touched_users,
                user_model=User,
            )
            auto_summary = {
                'updated_count': 0,
                'skipped_count': len(touched_users),
                'warnings': manager_sync_summary.get('warnings', []),
            }
        else:
            auto_summary = auto_apply_manager_chains(
                touched_users,
                fill_only_missing=False,
                commit=False,
                preserve_explicit_chain=True,
            )
            manager_sync_summary = sync_touched_users_manager_ids_from_sicils(
                touched_users,
                user_model=User,
            )
        db.session.commit()
        assignment_result = None
        from app.models import PerformancePeriod
        active_period = PerformancePeriod.query.filter_by(is_active=True).order_by(PerformancePeriod.id.desc()).first()
        if active_period:
            assignment_result = generate_assignments_for_active_period(period_id=active_period.id, actor_user_id=current_user.id)
        flash(
            f"İçe aktarma tamamlandı. Yeni: {created}, Güncellenen: {updated}, Atlanan admin: {skipped_admin}, kategori alanı işlendi, kaynak kurala göre yeniden kurulan zincir: {auto_summary.get('updated_count', 0)}",
            "success",
        )
        if has_explicit_manager_columns:
            flash('Excel içindeki yönetici sicil ve ikinci yönetici sicil alanları kaynak veri olarak korundu; amir seçili görünümü bu sicillere göre senkronlandı.', 'info')
        if assignment_result and assignment_result.get("ok"):
            flash(
                f"Aktif dönem görevleri de otomatik yenilendi. Oluşturulan/güncellenen görev: {assignment_result.get('created', 0)}, atlanan: {assignment_result.get('skipped', 0)}, gerçek uyarı: {assignment_result.get('warning_count', 0)}.",
                "success" if not assignment_result.get("warning_count") else "warning",
            )
        elif assignment_result and not assignment_result.get("ok"):
            flash(assignment_result.get("message", "Aktif dönem görevleri otomatik yenilenemedi."), "warning")
        if apply_ai_fixes and sum(int(v or 0) for v in autofix_stats.values()) > 0:
            flash(f"AI otomatik düzeltme akışı uygulandı. Başlık eşleştirme: {autofix_stats.get('header_alias_count', 0)}, metin temizliği: {autofix_stats.get('whitespace_trim_count', 0)}, rol çıkarımı: {autofix_stats.get('inferred_role_count', 0)}, varsayılan aktiflik: {autofix_stats.get('default_active_count', 0)}", "info")
        if auto_summary.get("warnings"):
            current_app.logger.warning(
                "Otomatik zincir ham uyarıları gizlendi: %s",
                " | ".join(auto_summary["warnings"][:10]),
            )
            flash(
                f"Otomatik zincir denetiminde {len(auto_summary.get('warnings', []))} kayıt için teknik not üretildi; ham liste kullanıcı ekranında gösterilmedi.",
                "info",
            )
        return redirect(url_for("main.admin_users"))

    return safe_render("excel_import.html", "<h3>Excel İçe Aktar</h3>", apply_ai_fixes=False, ai_excel_fix_panel=build_excel_fix_preview_ai_panel())
