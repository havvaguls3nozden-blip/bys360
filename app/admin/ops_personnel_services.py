"""Operational admin personnel services extracted from ops_routes.py.

Route decorators stay in ops_routes.py. This module contains implementation bodies.
"""
from __future__ import annotations

from io import BytesIO
from flask import flash, redirect, send_file, url_for
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from app.extensions import db
from app.models import User
from app.route_support import safe_render
from app.services.ai import build_personnel_profile_chain_ai_panel
from .ops_helpers import (
    _profile_full_name,
    _resolve_personnel_profile_hierarchy,
)

def download_personnel_template_impl():
    wb = Workbook()
    ws = wb.active
    ws.title = "Personel Şablonu"

    headers = [
        "Sicil No",
        "Ad",
        "Soyad",
        "E-Posta",
        "Unvan",
        "Birim",
        "Üst Birim",
        "Rol (opsiyonel)",
        "Yönetici Sicil (opsiyonel)",
        "İkinci Yönetici Sicil (opsiyonel)",
        "Üçüncü Yönetici Sicil (opsiyonel)",
    ]
    ws.append(headers)

    example_rows = [
        ["1001", "Ahmet", "Yılmaz", "ahmet.yilmaz@ktb.gov.tr", "Personel", "Eğitim ve Performans Çalışma Grubu", "Personel ve Destek Hizmetleri Grup Başkanlığı", "personel", "", "", ""],
        ["1002", "Ayşe", "Demir", "ayse.demir@ktb.gov.tr", "Koordinatör", "Eğitim ve Performans Çalışma Grubu", "Personel ve Destek Hizmetleri Grup Başkanlığı", "koordinator", "", "", ""],
    ]
    for row in example_rows:
        ws.append(row)

    header_fill = PatternFill("solid", fgColor="8B0000")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D1D5DB")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    widths = {
        "A": 16,
        "B": 18,
        "C": 18,
        "D": 30,
        "E": 24,
        "F": 34,
        "G": 34,
        "H": 20,
        "I": 20,
        "J": 24,
        "K": 24,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="bys360_personel_sablonu.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

def personnel_profile_impl(user_id: int):
    user = db.session.get(User, user_id)
    if not user:
        flash("Personel kaydı bulunamadı.", "danger")
        return redirect(url_for("main.personnel_list"))

    avatar_url = url_for("static", filename="img/default-avatar.svg")
    if hasattr(user, "profile_photo") and getattr(user, "profile_photo", None):
        avatar_url = url_for("static", filename=f"uploads/{user.profile_photo}")

    role_text = ""
    if hasattr(user, "role") and (getattr(user, "role", "") or "").strip():
        role_text = user.role
    elif hasattr(user, "role_label") and (getattr(user, "role_label", "") or "").strip():
        role_text = user.role_label

    resolved_hierarchy = _resolve_personnel_profile_hierarchy(user)

    profile_data = {
        "id": user.id,
        "full_name": _profile_full_name(user),
        "ad": getattr(user, "ad", "") or "-",
        "soyad": getattr(user, "soyad", "") or "-",
        "email": getattr(user, "email", "") or "-",
        "sicil_no": getattr(user, "sicil_no", "") or "-",
        "unvan": getattr(user, "unvan", "") or "-",
        "birim": resolved_hierarchy["birim"],
        "ust_birim": resolved_hierarchy["ust_birim"],
        "org_path": resolved_hierarchy["org_path"],
        "hierarchy_state": resolved_hierarchy["hierarchy_state"],
        "role_text": role_text or "-",
        "is_active": bool(getattr(user, "is_active", False)),
        "must_change_password": bool(getattr(user, "must_change_password", False)) if hasattr(user, "must_change_password") else False,
        "must_set_security_question": bool(getattr(user, "must_set_security_question", False)) if hasattr(user, "must_set_security_question") else False,
        "is_first_login": bool(getattr(user, "is_first_login", False)) if hasattr(user, "is_first_login") else False,
        "avatar_url": avatar_url,
        "manager_1": resolved_hierarchy["manager_1"],
        "manager_2": resolved_hierarchy["manager_2"],
        "manager_3": resolved_hierarchy["manager_3"],
        "yonetici_sicil": getattr(user, "yonetici_sicil", None),
        "ikinci_yonetici_sicil": getattr(user, "ikinci_yonetici_sicil", None),
        "ucuncu_yonetici_sicil": getattr(user, "ucuncu_yonetici_sicil", None),
    }

    from app.services.hr_operations_service import build_personnel_profile_hr_context

    return safe_render(
        "personnel_profile.html",
        "<h3>Personel Profili</h3>",
        profile=profile_data,
        ai_personnel_chain_panel=build_personnel_profile_chain_ai_panel(profile_data),
        hr_context=build_personnel_profile_hr_context(user),
    )
